"""As fórmulas do Excel exportado devem calcular exatamente o mesmo que o sistema (sem depender do Excel instalado)."""
import datetime as dt
import os
import pytest

from orcamento.importar_pt8 import carregar
from orcamento.otimizador import otimizar
from orcamento.calculo import verificar, media_rh, media_subitem
from orcamento.exportar import exportar
from orcamento.modelo import RubricaRH

pycel = pytest.importorskip('pycel')
DADOS = os.path.join(os.path.dirname(__file__), '..', '..', 'fase0', 'sejc', 'pt8_dados.json')


def test_formulas_do_excel_batem_com_o_sistema(tmp_path):
    novo = otimizar(carregar(DADOS), ajustar_horas=True)['projeto']
    f = str(tmp_path / 'grade.xlsx')
    exportar(novo, verificar(novo, {}, hoje=dt.date(2026, 9, 9)), {}, f)
    xl = pycel.ExcelCompiler(filename=f)
    esperadas = []
    for r in novo.rubricas:
        if isinstance(r, RubricaRH):
            esperadas.append(media_rh(r))
        else:
            esperadas.append(None); esperadas += [media_subitem(s) for s in r.subitens]
    for lin, esp in zip(range(4, 4 + len(esperadas)), esperadas):
        if esp is not None:
            assert round(xl.evaluate(f"'Grade Comparativa'!M{lin}") * 100) == esp, lin
    from openpyxl import load_workbook
    ws = load_workbook(f)['Plano de Aplicação']
    tot = max(r for r in range(1, ws.max_row + 1) if ws[f'B{r}'].value == 'Total')
    assert round(xl.evaluate(f"'Plano de Aplicação'!D{tot}") * 100) == novo.teto
    assert xl.evaluate(f"'Plano de Aplicação'!D{tot + 2}") == 'FECHA NO TETO'
