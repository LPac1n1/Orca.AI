"""Pedidos de 03 e 05/10/2026: descrição simples e sem marca, unidades com as letras certas, busca que entende o pedido (embalagem na frente,
"enlatada", abreviações, outros nomes do produto, títulos de vaga), valor no plano só o menor preço ou a média, comprovantes de CNPJ já
baixados, topo do projeto em todas as telas, menu, página inicial e o pacote para a Secretaria."""
import io
import os
import re
import zipfile

import pytest
from fastapi.testclient import TestClient


@pytest.fixture
def cliente(tmp_path, monkeypatch):
    monkeypatch.setenv('ORCAMENTO_DADOS', str(tmp_path))
    import importlib, orcamento.db as dbm
    monkeypatch.setattr(dbm, 'PASTA_DADOS', str(tmp_path))
    monkeypatch.setattr(dbm, 'PASTA_LOCAL', str(tmp_path / 'local'))
    from orcamento import comprovante_receita as CR
    import app as appmod
    importlib.reload(appmod)
    return TestClient(appmod.app)


def _novo(cliente):
    r = cliente.post('/projetos', data={'nome': 'T', 'teto': '100.000,00', 'cep': '01001-000'}, follow_redirects=False)
    return int(r.headers['location'].rsplit('/', 1)[-1])


# ---------------------------------------------------------------- valor no plano: o menor dos 3 preços ou a média, nunca o do meio
def test_valor_no_plano_e_o_menor_preco_ou_a_media(cliente):
    from orcamento import db
    from orcamento.calculo import verificar
    from orcamento.modelo import Subitem
    from orcamento.otimizador import otimizar
    from orcamento.produtos import teto
    from orcamento.regras import valores_do_plano
    assert valores_do_plano([1000, 1200, 2000]) == [1000, 1400] and valores_do_plano([500, 500, 500]) == [500]
    assert teto.dominio([1000, 1200, 2000]) == (1400, [1000, 1400])                    # a pesquisa da rubrica também não escolhe o preço do meio
    pid = _novo(cliente)
    cliente.post(f'/p/{pid}/rubrica', data={'tipo': 'material', 'nome': 'Alimentação', 'regra': 'mercado', 'meses': '1'})
    p = db.carregar(pid)[0]
    p.rubricas[0].subitens = [Subitem(descricao='Arroz', qtd=1, precos=[1000, 1200, 2000], valor_plano=1200)]
    assert [a.regra for a in verificar(p) if a.regra == 'S01'] == ['S01']             # 12,00 é o preço do meio: a verificação aponta
    for teto_do_projeto, esperado in ((1000, 1000), (1400, 1400)):
        p.teto = teto_do_projeto
        r = otimizar(p)
        assert r['status'] == 'OK' and r['projeto'].rubricas[0].subitens[0].valor_plano == esperado
        assert not [a for a in verificar(r['projeto']) if a.regra == 'S01']
    p.teto = 1200                                                                       # só fecharia com o preço do meio: não fecha
    assert otimizar(p)['status'] == 'SEM_SOLUCAO'


# ---------------------------------------------------------------- unidades de medida
def test_unidades_de_medida_com_as_letras_certas():
    from orcamento.produtos.texto import formatar_medidas as f
    casos = {'1l': '1L', '500ML': '500mL', '500 ml': '500 mL', '1Kg': '1kg', '5 KG': '5 kg', '2 litros': '2L', '1,35l': '1,35L', '30CM': '30cm', '100 gramas': '100g',
             '110v': '110V', '50L 50 Unidades': '50L 50 Unidades', '5000 Unidades': '5000 Unidades', '4G': '4G', 'Tamanho M': 'Tamanho M', '1 lata': '1 lata',
             '210x297MM': '210x297mm', '75g/m²': '75g/m²', '12 UN': '12 un', '1.5 LT': '1.5 L', '26/6': '26/6'}
    for antes, depois in casos.items():
        assert f(antes) == depois, antes
    assert f(f('2 litros de suco 500ML')) == '2L de suco 500mL' and f(None) is None and f('') == ''


# ---------------------------------------------------------------- descrição simples, sem marca
REAIS = [  # (anúncios das 3 lojas, pedido, marca, nome simples esperado, especificação esperada)
    (['Manteiga com Sal Aviação 500g', 'Manteiga de Primeira Qualidade com Sal Aviação Pote 500g', 'Manteiga com Sal Aviação Pote 500g'], 'Manteiga com Sal', 'Aviação',
     'Manteiga com Sal', '500g'),
    (['Pão de Forma Artesano Pullman Pacote 500g', 'Pão de Forma Artesano Pullman 500g', 'Pão de Forma Original Pullman Artesano Pacote 500g'], 'Pão de Forma', 'Pullman',
     'Pão de Forma', '500g'),
    (['Pão de Forma Wickbold Integral 450g', 'Pão de Forma Integral Premium Wickbold Pacote 450g', 'Pão de Forma Integral Wickbold Viva Integralmente Premium Pacote 450g'],
     'Pão de Forma Integral', 'Wickbold', 'Pão de Forma Integral', '450g'),
    (['Suco Maguary Seleção Uva e Maçã 1,35L', 'Suco Seleção 100% Uva E Maçã Maguary 1,35L', 'Suco Misto Uva E Maçã Maguary Garrafa 1,35l'], 'Suco de Maça', 'Maguary',
     'Suco de Uva e Maçã', '1,35L'),
    (['Bloco Adesivo Post-it Cubo Ultra 400 Folhas 47.6mm x 47.6mm - 3M', 'Bloco de Notas Adesivo Post-it Cubo Ultra 47,6x47,6mm 400 FL Cores Sor',
      'Bloco Post-It Cubo 400 Folhas Ultra'], 'Bloco de Notas', 'Post-it', 'Bloco Adesivo', '47.6mm 400 Folhas'),
    (['Régua Dello Poliestireno Cristal 30cm 3109 H.0100 1 UN', 'Régua 30cm Cristal - Dello', 'Régua 30cm Cristal'], 'Pasta', None, 'Régua', '30cm Cristal'),
    (['Lápis de Cor Bic Evolution 12 cores un', 'Lápis de cor 12 cores Evolution + 4 lápis grafite - Bic', 'Lápis De Cor Evolution Color Grátis 4 Lápis - Bic'],
     'Lápis Grafite', 'Bic', 'Lápis de Cor', '12 Cores'),
    (['Limpador Multiuso Limpol Essence 500ml', 'Limpador Multiuso Limpol 500ml 1 UN', 'Limpador Multiuso Limpol Essence 500ml'], 'Sabão em Pó', 'Limpol',
     'Limpador Multiuso', '500mL'),
    (['Suco Concentrado Maracujá Maguary Garrafa 1l', 'Suco Concentrado Maguary Maracujá 1L', 'Suco Concentrado de Maracujá Maguary 1L'], None, 'Maguary',
     'Suco Concentrado de Maracujá', '1L'),
    (['Café Torrado e Moído Tradicional Pilão Pacote 500g', 'Café em Pó Pilão Tradicional Promoção 500g 1 UN', 'Café Torrado e Moído Tradicional Pilão Pacote 500g'],
     'Café', 'Pilão', 'Café', '500g'),
]


@pytest.mark.parametrize('titulos,pedido,marca,nome,esp', REAIS)
def test_nome_simples_do_produto_sem_marca_linha_e_medida(titulos, pedido, marca, nome, esp):
    from orcamento.produtos import texto as T
    achado = T.nome_simples(titulos, pedido, marca)
    assert achado == nome and T.especificacao_do_produto(titulos, achado) == esp
    palavras = {T.sa(w) for t in titulos + [pedido or ''] for w in re.findall(r'\w+', t)}
    assert all(T.sa(w) in palavras for w in achado.split())                              # nenhuma palavra inventada
    assert not T.precisa_arrumar(achado, marca, trocado=True)                            # o nome já é simples: não é mexido de novo


def test_linha_do_fabricante_nao_e_marca(cliente):
    """Decisão de 05/10/2026: "Evolution" é uma linha de lápis da Bic — a marca do item é o fabricante. A pesquisa nova já acha "Bic", e o item
    gravado com "Evolution" é acertado sozinho quando o projeto abre (os 3 anúncios precisam dizer o fabricante)."""
    from orcamento import db, servico
    from orcamento.modelo import RubricaMaterial, Subitem, descricao_completa
    from orcamento.produtos import texto as T, identidade as ID
    anuncios = ['Lápis de Cor Bic Evolution 12 cores un', 'Lápis de cor 12 cores Evolution + 4 lápis grafite - Bic', 'Lápis De Cor Evolution Color Grátis 4 Lápis - Bic']
    assert 'evolution' not in ID.MARCAS_N and [ID.marca_de(dict(nome=a)) for a in anuncios] == ['bic'] * 3
    assert servico.marca_do_produto(dict(ofertas=[dict(nome=a) for a in anuncios])) == 'Bic'
    assert T.fabricante_da_linha('Evolution', anuncios) == 'Bic' and T.fabricante_da_linha('evolution', ['LÁPIS BIC EVOLUTION', 'Lápis BIC Evolution']) == 'Bic'
    assert T.fabricante_da_linha('Evolution', anuncios[:1] + ['Lápis de cor Evolution 12 cores']) is None      # um anúncio sem o fabricante: não mexe
    assert T.fabricante_da_linha('Pilot', anuncios) is None and T.fabricante_da_linha(None, anuncios) is None and T.fabricante_da_linha('Evolution', []) is None
    pid = _novo(cliente)
    p = db.carregar(pid)[0]
    p.rubricas = [RubricaMaterial(item=1, descricao='Material', meses=10, subitens=[
        Subitem(descricao='Lápis de Cor', marca='Evolution', especificacao='12 Cores', qtd=1, precos=[900, 990, 1100], valor_plano=990, produtos=anuncios, nivel=2,
                descricao_original='Lápis Grafite', marca_original='', especificacao_original=''),
        Subitem(descricao='Caneta Esferográfica', marca='Pilot', especificacao='Azul', qtd=1, precos=[300, 319, 350], valor_plano=319,
                produtos=['Caneta Esferográfica Pilot Azul', 'Caneta esferográfica azul - Pilot', 'Caneta Esferográfica Pilot BP-1RT Azul'])])]
    db.salvar(pid, p)
    cliente.get(f'/p/{pid}')                                                                                  # abrir o projeto acerta o que estava gravado
    p = db.carregar(pid)[0]
    assert [(s.descricao, s.marca, s.especificacao) for s in p.rubricas[0].subitens] == [('Lápis de Cor', 'Bic', '12 Cores'), ('Caneta Esferográfica', 'Pilot', 'Azul')]
    assert descricao_completa(p.rubricas[0].subitens[0]) == 'Lápis de Cor Bic 12 Cores' and p.rubricas[0].subitens[0].descricao_original == 'Lápis Grafite'
    v = db.carregar(pid)[1]
    cliente.get(f'/p/{pid}')
    assert db.carregar(pid)[1] == v                                                                           # e não mexe de novo


def test_item_digitado_marca_vai_para_o_campo_marca(cliente):
    from orcamento import db
    from orcamento.modelo import pedido_do_subitem
    from orcamento.produtos import texto as T
    assert T.arrumar_item('Suco de Uva Del Valle', None, '1l') == ('Suco de Uva', 'Del Valle', '1L')
    assert T.arrumar_item('Manteiga com Sal Aviação', 'Aviação', '500G') == ('Manteiga com Sal', 'Aviação', '500g')
    for intocado in (('Água de Coco', None, '1L'), ('Pão da Casa', None, None), ('Papel Brilhante A4', None, None), ('Cândida', None, '2L'), ('Bloco Post-it', None, None)):
        assert T.arrumar_item(*intocado) == intocado                                     # "coco", "da" e palavras comuns não são marca; a descrição nunca fica vazia
    pid = _novo(cliente)
    cliente.post(f'/p/{pid}/rubrica', data={'tipo': 'material', 'nome': 'Alimentação', 'regra': 'mercado', 'meses': '10'})
    cliente.post(f'/p/{pid}/mat/1', data={'descricao': 'Alimentação', 'meses': '10', 'n_sub': '0', 'ndesc': ['Café Pilão', 'Leite'], 'nmarca': ['', ''],
                                           'nesp': ['500G', '1l'], 'nqtd': ['2', '3']})
    subs = db.carregar(pid)[0].rubricas[0].subitens
    assert [(s.descricao, s.marca, s.especificacao) for s in subs] == [('Café', 'Pilão', '500g'), ('Leite', None, '1L')]
    assert pedido_do_subitem(subs[0]) == 'Café Pilão 500g'


def test_produto_trocado_fica_com_nome_simples_e_itens_antigos_sao_acertados(cliente):
    from orcamento import db, servico
    from orcamento.modelo import Subitem, pedido_do_subitem, descricao_completa
    # troca pela categoria: o tipo é o do catálogo, a medida vai para a especificação
    s = Subitem(descricao='Sabão em Pó', especificacao='1,6Kg', qtd=3)
    servico.campos_do_produto(s, dict(nivel=3, substituto='Limpador multiuso 500mL', ofertas=[dict(loja=l, nome=n, ean=None) for l, n in
                                      (('a', 'Limpador Multiuso Limpol Essence 500ml'), ('b', 'Limpador Multiuso Limpol 500ml 1 UN'), ('c', 'Limpador Multiuso Limpol Essence 500ml'))]))
    assert (s.descricao, s.marca, s.especificacao, pedido_do_subitem(s)) == ('Limpador Multiuso', 'Limpol', '500mL', 'Sabão em Pó 1,6Kg')
    # itens gravados antes: a tela abre com o acerto feito e gravado (uma vez), e os pontos revisados acompanham o nome novo
    pid = _novo(cliente)
    cliente.post(f'/p/{pid}/rubrica', data={'tipo': 'material', 'nome': 'Alimentação', 'regra': 'mercado', 'meses': '10'})
    p = db.carregar(pid)[0]
    r = p.rubricas[0]
    r.subitens = [Subitem(descricao='Manteiga de Primeira Qualidade com Sal Aviação Pote', marca='Aviação', especificacao='500g', descricao_original='Manteiga com Sal',
                          marca_original='', especificacao_original='1Kg', nivel=1, qtd=2,
                          produtos=['Manteiga com Sal Aviação 500g', 'Manteiga de Primeira Qualidade com Sal Aviação Pote 500g', 'Manteiga com Sal Aviação Pote 500g']),
                  Subitem(descricao='Leite', especificacao='1l', qtd=1), Subitem(descricao='Café', marca='Pilão', especificacao='500g', qtd=1)]
    velho = f'R10|Item 1 – Alimentação / {r.subitens[0].descricao}|faltam preços'
    p.revisados[velho] = dict(em='x')
    v0 = db.salvar(pid, p)
    pag = cliente.get(f'/p/{pid}/mat/1').text
    assert 'name="sdesc0" type="text" value="Manteiga com Sal"' in pag and 'Primeira Qualidade' not in re.sub(r'<summary class="detalhe.*?</details>', '', pag, flags=re.S).split('Detalhes e comprovante')[0]
    p, v = db.carregar(pid)
    s = p.rubricas[0].subitens
    assert v == v0 + 1 and (s[0].descricao, s[0].marca, s[0].especificacao) == ('Manteiga com Sal', 'Aviação', '500g') and pedido_do_subitem(s[0]) == 'Manteiga com Sal 1Kg'
    assert descricao_completa(s[0]) == 'Manteiga com Sal Aviação 500g' and s[1].especificacao == '1L' and (s[2].descricao, s[2].marca) == ('Café', 'Pilão')
    assert list(p.revisados) == ['R10|Item 1 – Alimentação / Manteiga com Sal|faltam preços']
    cliente.get(f'/p/{pid}/mat/1'); cliente.get(f'/p/{pid}')
    assert db.carregar(pid)[1] == v                                                       # o acerto é feito uma vez só


# ---------------------------------------------------------------- busca que entende o pedido
def test_pedido_com_a_embalagem_na_frente_ou_como_adjetivo():
    from orcamento.produtos import identidade as ID
    n = ID.normalizar_pedido
    assert n('Caixa Caneta Esferográfica Azul') == 'Caneta Esferográfica Azul caixa' and n('Caixa com 50 canetas azuis') == 'canetas azuis caixa com 50'
    assert n('Sardinha Enlatada') == 'Sardinha em lata' and n('Lata de Sardinha') == 'Sardinha lata' and n('Clips cx c/ 100') == 'Clips caixa com 100'
    for proprio in ('Caixa Organizadora 30L', 'Garrafa Térmica 1L', 'Saco de Lixo 50L', 'Kit Escolar', 'Leite caixa 1L', 'Sardinha em Lata'):
        assert n(proprio) == proprio and n(n(proprio)) == proprio                          # a embalagem é o próprio produto, ou já está no lugar
    assert n(n('Caixa de Canetas Azuis')) == n('Caixa de Canetas Azuis') and ID.familia_padrao('Caixa Caneta Esferográfica Azul') == 'caneta esferografica'
    caixa, avulsa = 'Caneta Esferográfica Bic Cristal Azul Caixa com 50 Unidades', 'Caneta Esferográfica Bic Cristal Azul 1 UN'
    for pedido in ('Caixa Caneta Esferográfica Azul', 'Caixa de Canetas Esferográficas Azuis'):
        assert ID.nivel(pedido, caixa) == 0 and ID.nivel(pedido, 'Caneta Esferográfica Bic Cristal Preta Caixa com 50') is None   # outra cor não atende
    grupo = lambda nome: dict(por_loja={l: dict(nome=nome, preco=100) for l in 'abc'})
    d_caixa, d_avulsa = (ID.distancia('Caixa Caneta Esferográfica Azul', grupo(x), 0, 'abc') for x in (caixa, avulsa))
    assert d_caixa[0] < d_avulsa[0] and d_avulsa[1] == ['caixa'] and d_caixa[1] == []      # a caixa vem antes da caneta avulsa
    assert ID.distancia('Caneta Esferográfica Azul', grupo(avulsa), 0, 'abc')[1] == []      # sem "caixa" no pedido, a avulsa é o pedido
    for pedido in ('Sardinha Enlatada', 'Sardinha em Lata', 'Lata de Sardinha', 'Sardinha'):
        assert ID.nivel(pedido, 'Sardinha Coqueiro com Óleo 125g') == 0                    # as lojas não escrevem "lata": não é exigido
    assert ID.nivel('Sardinha em Lata', 'Sardinha Coqueiro Pouch 100g') != 0               # mas outra embalagem citada não é o pedido
    assert ID.nivel('Sabonete Líquido 500mL', 'Sabonete Líquido Refil 500ml') != 0 and ID.nivel('Sabonete Líquido Refil 500mL', 'Sabonete Líquido 500ml') != 0
    assert ID.nivel('Canetas Azuis', 'Caneta Esferográfica Azul Bic') == 0 and ID.nivel('Caneta Preta', 'Caneta Esferográfica Preto Bic') == 0   # plural e gênero
    assert [ID.marca_de(dict(nome=x)) for x in ('Água de Coco Sococo 1L', 'Água de Coco Kero Coco 1L', 'Pão da Casa 500g', 'Limpador Mr. Músculo 500ml')] == \
        ['sococo', 'kero coco', None, 'mr musculo']                                        # "coco" e "da" não são marca; a de duas palavras é achada


def test_ia_entende_o_pedido_mas_as_regras_conferem(monkeypatch):
    from orcamento import ia
    from orcamento.produtos import cesta, identidade as ID
    resposta = {'Lápis Grafite': dict(produto='lápis grafite', sinonimos=['lápis preto', 'lápis hb', 'Lápis Faber-Castell', 'borracha branca'],
                                      buscas=['lápis grafite hb', 'lápis preto', 'lápis bic'], varias_unidades=False),
                'Borracha Branca': dict(produto='apontador', sinonimos=['apontador'], buscas=[], varias_unidades=None)}
    monkeypatch.setattr(ia, 'disponivel', lambda: True)
    monkeypatch.setattr(ia, 'entender_pedidos', lambda pedidos, pid=None: resposta)
    m = cesta.Motor([dict(desc='Lápis Grafite', qtd=1), dict(desc='Borracha Branca', qtd=1)], ['kalunga', 'gimba', 'lepok'], '01001-000')
    assert ID.nivel('Lápis Grafite', 'Lápis HB Nº2 Sextavado Leo&Leo') != 0
    m._entender()
    e = m.entendido['Lápis Grafite']
    assert e['nomes'] == ['lápis preto', 'lápis hb']                                       # nome com marca e nome de OUTRO item da rubrica ficam de fora
    assert m.buscas['Lápis Grafite'] == ['Lápis Grafite', 'lápis grafite hb', 'lápis preto']   # busca com marca que o pedido não cita não entra
    assert ID.nivel('Lápis Grafite', 'Lápis HB Nº2 Sextavado Leo&Leo') == 0                 # o outro nome do mesmo produto passa a valer
    assert 'nomes' not in m.entendido['Borracha Branca'] and ID.nivel('Borracha Branca', 'Apontador com Depósito') is None   # "tipo" que não está no pedido é ignorado
    ID._DINAMICOS.clear(); ID._VARIAS.clear()


def test_titulo_de_vaga_entendido_sem_afrouxar_o_cargo():
    from orcamento import vagas as V
    sim = [('Aux. Administrativo', 'Auxiliar Administrativo'), ('Auxiliar Administrativo - Zona Sul', 'Auxiliar Administrativo'), ('Vaga de Assistente Social', 'Assistente Social'),
           ('Assistente Social (Temporário)', 'Assistente Social'), ('Auxiliar de Serviços Gerais PCD', 'Auxiliar de Serviços Gerais'), ('Aux. Serv. Ger.', 'Auxiliar de Serviços Gerais'),
           ('Coordenador(a) de Projetos | São Paulo', 'Coordenador de Projetos'), ('Psicóloga - SP', 'Psicólogo'), ('Coord. Pedagógico', 'Coordenador Pedagógico'),
           # o cargo seguido de uma extensão qualquer (decisão da OSC, 06/10/2026)
           ('Assistente Social - Hospitalar', 'Assistente Social'), ('Psicólogo - Coca-Cola', 'Psicólogo'), ('Orientador Socioeducativo - Educação', 'Orientador Socioeducativo'),
           ('Coordenador de Projetos: Educação', 'Coordenador de Projetos'), ('Psicólogo (Clínica Vida)', 'Psicólogo'), ('Coordenador de Projetos|Campinas', 'Coordenador de Projetos')]
    for titulo, cargo in sim:
        assert V.titulo_exato(titulo, cargo, 'São Paulo', 'SP'), titulo
    nao = [('Psicólogo Clínico', 'Psicólogo'), ('Assistente Social Júnior', 'Assistente Social'), ('Auxiliar Administrativo II', 'Auxiliar Administrativo'),
           ('Coordenador de Projetos Sociais', 'Coordenador de Projetos'), ('Supervisor Administrativo', 'Auxiliar Administrativo'),
           ('Estágio em Psicologia', 'Psicólogo'), ('Psicólogo Clínico - Hospital', 'Psicólogo'), ('Hospital - Psicólogo', 'Psicólogo'),
           ('Assistente do Coordenador de Projetos', 'Coordenador de Projetos')]
    for titulo, cargo in nao:
        assert not V.titulo_exato(titulo, cargo, 'São Paulo', 'SP'), titulo
    assert V.cargo_para_busca('Orientador(a) Socioeducativo (recibo ou MEI) - 20h semanais') == 'Orientador Socioeducativo'
    assert V.cargo_para_busca('Aux. de Limpeza') == 'auxiliar de Limpeza' and V.cargo_para_busca('Psicólogo') == 'Psicólogo'
    assert 'vagas-de-emprego-orientador-socioeducativo.aspx' in V.fontes('Orientador(a) Socioeducativo (recibo ou MEI) - 20h semanais')[0][1]
    assert V.titulo_confere('Aux. Administrativo', 'Auxiliar Administrativo') and V.chave_cargo('Auxiliar de Serviços Gerais') == 'auxiliar de servico geral'


# ---------------------------------------------------------------- telas
def test_topo_do_projeto_em_todas_as_telas_menu_e_pagina_inicial(cliente):
    pid = _novo(cliente)
    for rota, atual in ((f'/p/{pid}', None), (f'/p/{pid}/cnpjs', 'CNPJs e comprovantes'), (f'/p/{pid}/historico', 'Histórico')):
        pag = cliente.get(rota).text
        assert pag.count('<h1') == 1 and 'Teto do projeto' in pag and 'Diferença para o teto' in pag and 'class="subnav"' in pag, rota
        assert f'href="/p/{pid}/pacote"' in pag and f'href="/p/{pid}/planilhas"' in pag, rota
        menu = re.search(r'<nav class="subnav".*?</nav>', pag, re.S).group(0)
        ordem = [menu.index(x) for x in ('Visão geral', 'Mão de obra', 'Materiais e serviços', 'Verificação', 'CNPJs e comprovantes', 'Configuração', 'Histórico')]
        assert ordem == sorted(ordem), rota                                                 # Configuração é a penúltima; Histórico, a última
        if atual:   # fora da tela do projeto, os botões levam à parte certa dele e a tela atual fica marcada
            assert f'href="/p/{pid}#configuracao"' in menu and 'data-secoes' not in menu and re.search(r'aria-current="true">.{0,200}' + atual, menu, re.S)
        else:
            assert 'href="#configuracao"' in menu and 'data-secoes' in menu
    inicio = cliente.get('/').text
    assert 'style="max-width: 760px"' not in inicio and 'class="formulario formulario--projeto"' in inicio   # o formulário ocupa a largura toda


def test_comprovante_ja_baixado_aparece_antes_de_importar(cliente, monkeypatch, tmp_path):
    import pymupdf
    from orcamento import db, comprovante_receita as CR
    from orcamento.modelo import PesquisaSalarial
    pid = _novo(cliente)
    cliente.post(f'/p/{pid}/rubrica', data={'tipo': 'rh', 'nome': 'Psicólogo', 'horas_mes': '100', 'meses': '10'})
    p = db.carregar(pid)[0]
    p.rubricas[0].pesquisas = [PesquisaSalarial(nome='ALFA', cnpj='11.222.333/0001-81', valor=200000), PesquisaSalarial(nome='BETA', cnpj='11.444.777/0001-61', valor=200000)]
    db.salvar(pid, p)
    assert cliente.get(f'/p/{pid}/cnpjs').text.count('Falta emitir</span>') == 2 and cliente.get(f'/p/{pid}/cnpjs/baixados').json() == {}
    doc = pymupdf.open(); pg = doc.new_page()
    pg.insert_text((40, 60), 'COMPROVANTE DE INSCRIÇÃO E DE SITUAÇÃO CADASTRAL\nNÚMERO DE INSCRIÇÃO 11.222.333/0001-81\nNOME EMPRESARIAL ALFA LTDA TÍTULO DO ESTABELECIMENTO\n'
                             'SITUAÇÃO CADASTRAL ATIVA\nEmitido no dia 05/10/2026 às 10:00:00')
    arq = os.path.join(CR.pasta_dos_comprovantes(), 'comprovante alfa.pdf')
    doc.save(str(arq))
    assert CR.pasta_dos_comprovantes() == os.environ['ORCAMENTO_COMPROVANTES'] and CR.pdfs_recentes() == [str(arq)]   # o PDF está na pasta dos comprovantes
    assert cliente.get(f'/p/{pid}/cnpjs/baixados').json() == {'11.222.333/0001-81': dict(arquivo='comprovante alfa.pdf', situacao='ATIVA', emitido_em='2026-10-05T10:00:00', importado=False)}
    pag = cliente.get(f'/p/{pid}/cnpjs').text
    assert 'Já baixado: falta importar' in pag and 'comprovante alfa.pdf' in pag and pag.count('Falta emitir</span>') == 1   # dá para ver qual falta, antes de importar


# ---------------------------------------------------------------- pacote para a Secretaria
def _pdf(texto):
    import pymupdf
    d = pymupdf.open(); d.new_page().insert_text((40, 60), texto)
    return d.tobytes()


def test_pacote_para_a_secretaria_com_planilhas_e_pdfs_juntados(cliente):
    import pymupdf
    from openpyxl import load_workbook
    from orcamento import db
    from orcamento.modelo import PesquisaSalarial, Evidencia, Subitem, Fonte
    pid = _novo(cliente)
    cliente.post(f'/p/{pid}/rubrica', data={'tipo': 'rh', 'nome': 'Psicólogo', 'horas_mes': '100', 'meses': '10'})
    cliente.post(f'/p/{pid}/rubrica', data={'tipo': 'material', 'nome': 'Alimentação', 'regra': 'mercado', 'meses': '10'})
    p = db.carregar(pid)[0]
    guardar = lambda nome, texto: db.guardar_arquivo(pid, nome, _pdf(texto))[1]
    p.rubricas[0].pesquisas = [PesquisaSalarial(nome='ALFA LTDA', cnpj='11.222.333/0001-81', valor=200000, evidencia=Evidencia(arquivo=guardar('vaga1.pdf', 'PAGINA DA VAGA ALFA'))),
                               PesquisaSalarial(nome='BETA: S/A', cnpj='11.444.777/0001-61', valor=220000, evidencia=Evidencia(arquivo=guardar('vaga2.pdf', 'PAGINA DA VAGA BETA'))),
                               PesquisaSalarial(nome='GAMA', cnpj='', valor=240000)]
    p.rubricas[0].valor_mensal_plano = 100000
    fontes = lambda t: [Fonte(nome=f'LOJA {k}', plataforma=f'Loja{k}', cnpj='11.222.333/0001-81', evidencia=Evidencia(arquivo=guardar(f'{t}{k}.pdf', f'PRODUTO {t} LOJA {k}'))) for k in (1, 2, 3)]
    p.rubricas[1].subitens = [Subitem(descricao='Café', marca='Pilão', especificacao='500g', qtd=2, precos=[1000, 1200, 2000], valor_plano=1400, fontes=fontes('cafe')),
                              Subitem(descricao='Leite', especificacao='1L', qtd=3, precos=[500, 600, 700], valor_plano=500, fontes=fontes('leite'))]
    p.comprovantes_cnpj['11.222.333/0001-81'] = dict(arquivo=guardar('cnpj_alfa.pdf', 'COMPROVANTE DE INSCRICAO ALFA'), situacao='ATIVA', emitido_em='2026-10-05')
    db.salvar(pid, p)
    r = cliente.get(f'/p/{pid}/pacote')
    assert r.status_code == 200 and r.headers['content-type'] == 'application/zip'
    z = zipfile.ZipFile(io.BytesIO(r.content))
    nomes = [n.split('/', 1)[1] for n in z.namelist()]
    assert nomes[:2] == ['Plano de Aplicação e Comparativo de Preço.xlsx', 'Plano de Aplicação e Comparativo de Preço.pdf'] and nomes[-1] == 'LEIA-ME.txt'
    assert nomes[2:-1] == ['Orçamentos/01. Psicólogo/1. ALFA LTDA.pdf', 'Orçamentos/01. Psicólogo/2. BETA S A.pdf'] + \
        [f'Orçamentos/02. Alimentação/{i}/{k}. Loja{k}.pdf' for i in ('01. Café Pilão 500g', '02. Leite 1L') for k in (1, 2, 3)]   # por rubrica e item, na ordem do plano
    paginas = lambda n: [pg.get_text().strip() for pg in pymupdf.open(stream=z.read([x for x in z.namelist() if x.endswith(n)][0]), filetype='pdf')]
    assert paginas('1. ALFA LTDA.pdf') == ['PAGINA DA VAGA ALFA', 'COMPROVANTE DE INSCRICAO ALFA']       # a vaga e, no mesmo PDF, o comprovante de CNPJ da empresa
    assert paginas('01. Café Pilão 500g/2. Loja2.pdf') == ['PRODUTO cafe LOJA 2', 'COMPROVANTE DE INSCRICAO ALFA'] and paginas('2. BETA S A.pdf') == ['PAGINA DA VAGA BETA']
    leia = z.read([x for x in z.namelist() if x.endswith('LEIA-ME.txt')][0]).decode('utf-8')
    assert 'Sem PDF da pesquisa (não entrou no pacote): item 1 (Psicólogo), pesquisa 3 — GAMA' in leia and 'Sem o comprovante de CNPJ junto: item 1 (Psicólogo), pesquisa 2 — BETA: S/A (11.444.777/0001-61)' in leia
    # as planilhas, na formatação da planilha de pré-cálculos
    wb = load_workbook(io.BytesIO(cliente.get(f'/p/{pid}/planilhas').content))
    assert wb.sheetnames == ['Plano de Aplicação', 'Cronograma fisico-financeiro', 'Etapa e Fases', 'Cronograma de desembolso', 'Comparativo de Preço']   # as abas da planilha de pré-cálculos
    c = wb['Comparativo de Preço']
    assert c['A1'].value == 'COMPARATIVO DE PREÇOS' and (c['A1'].font.name, c['A1'].font.sz, c['A1'].font.b) == ('Verdana', 8, True) and c['B5'].font.name == 'Times New Roman'
    assert {str(m) for m in c.merged_cells.ranges} >= {'A1:M1', 'A2:A4', 'B2:B4', 'C2:C4', 'D2:F2', 'D3:D4', 'E3:E4', 'G2:I2', 'J2:L2', 'M2:M4', 'A6:A8'}
    assert [c[x].value for x in ('A2', 'B2', 'C2', 'D2', 'D3', 'E3', 'F3', 'F4', 'M2')] == ['ITEM', 'DESCRIÇÃO DO ITEM', 'Q\nT\nD\nE', 'ORÇAMENTO 1', 'PREÇO UNITÁRIO', 'PREÇO TOTAL ', 'CNPJ', 'FORNECEDOR', 'Valor Médio']
    assert [c[x].value for x in ('A5', 'B5', 'C5', 'D5', 'E5', 'F5', 'M5')] == [1, 'Psicólogo (recibo ou MEI) - 100h mensal', 1, 2000, '=D5*C5', 'ALFA LTDA - CNPJ: 11.222.333/0001-81', '=ROUND((D5+G5+J5)/3,2)']
    assert [c[x].value for x in ('A6', 'B6', 'E6', 'M6', 'B7', 'C7', 'D7', 'M7')] == [2, 'Alimentação', '=SUM(E7:E8)', '=ROUND((E6+H6+K6)/3,2)', 'Café Pilão 500g', 2, 10, '=ROUND((D7+G7+J7)/3,2)']
    assert c['D5'].number_format.startswith('_-"R$ "* #,##0.00') and c['D5'].alignment.wrap_text and c.column_dimensions['F'].width == 13.3
    assert (c.print_title_rows, c.page_setup.orientation, c.page_setup.scale) == ('$1:$4', 'landscape', 74) and c.print_area.endswith('$A$1:$M$8')
    pl = wb['Plano de Aplicação']
    assert [pl[x].value for x in ('C1', 'C2', 'D2', 'E2', 'F2', 'G2', 'H2', 'D3')] == ['PLANO DE APLICAÇÃO', 'Item', 'Descrição', 'Valor Unitário', 'Valor Total', 'Concedente\n (SJC)',
                                                                                       'Proponente\n (entidade)', 'Recursos Humanos']
    assert [pl[x].value for x in ('C4', 'D4', 'E4', 'F4', 'C5', 'D5', 'E5', 'F5', 'D6', 'E6', 'C8', 'F8')] == \
        [1, 'Psicólogo (recibo ou MEI) - 100h mensal', 1000, '=ROUND(E4*10,2)', 2, 'Alimentação', '=SUM(E6:E7)', '=ROUND(E5*10,2)', 'Café Pilão 500g (2 unidades x R$14,00)', '=2*14.0', 'Total', '=F4+F5']
    assert pl['C1'].font.name == 'Aptos Narrow' and {str(m) for m in pl.merged_cells.ranges} >= {'C1:H1', 'C5:C7', 'F5:F7', 'G5:G7', 'C8:E8'} and pl.print_area.endswith('$C$1:$H$8')
    # cada fórmula leva junto o valor calculado: quem abre sem recalcular (Modo de Exibição Protegido do Excel, pré-visualização, celular) vê os números do sistema
    sem = load_workbook(io.BytesIO(cliente.get(f'/p/{pid}/planilhas').content), data_only=True)
    assert [sem['Plano de Aplicação'][x].value for x in ('E3', 'F3', 'F4', 'G4', 'E5', 'F5', 'G5', 'E6', 'E7', 'F8', 'G8')] == [1000, 10000, 10000, 10000, 43, 430, 430, 28, 15, 10430, 10430]
    assert [sem['Comparativo de Preço'][x].value for x in ('E5', 'H5', 'K5', 'M5', 'D6', 'E6', 'H6', 'K6', 'M6', 'E7', 'M7', 'E8', 'M8')] ==         [2000, 2200, 2400, 2200, 15, 35, 42, 61, 46, 20, 14, 15, 6]
    for ws in wb.worksheets:   # nenhuma fórmula sem valor
        for linha_ in ws.iter_rows():
            for c in linha_:
                if isinstance(c.value, str) and c.value.startswith('='):
                    assert sem[ws.title][c.coordinate].value is not None, (ws.title, c.coordinate)


def test_pasta_dos_comprovantes_e_orca_ai_em_documentos(cliente, monkeypatch, tmp_path):
    """Decisão da OSC (05/10/2026): os comprovantes da Receita são salvos e procurados na pasta "Orça.AI", em Documentos — não mais em Downloads."""
    from orcamento import comprovante_receita as CR
    monkeypatch.delenv('ORCAMENTO_COMPROVANTES')
    docs = tmp_path / 'OneDrive' / 'Documentos'
    docs.mkdir(parents=True)
    monkeypatch.setattr(CR, 'pasta_documentos', lambda: str(docs))                          # a pasta Documentos que o Windows aponta (pode estar no OneDrive)
    assert CR.pasta_dos_comprovantes() == str(docs / 'Orça.AI') and not (docs / 'Orça.AI').exists()
    assert CR.pasta_dos_comprovantes(criar=True) == str(docs / 'Orça.AI') and (docs / 'Orça.AI').is_dir()
    (docs / 'Orça.AI' / 'a.pdf').write_bytes(b'%PDF-1.4'); (docs / 'Orça.AI' / 'nota.txt').write_text('x')
    (tmp_path / 'Downloads').mkdir(); (tmp_path / 'Downloads' / 'b.pdf').write_bytes(b'%PDF-1.4')
    assert CR.pdfs_recentes() == [str(docs / 'Orça.AI' / 'a.pdf')] and CR.baixados() == {}   # só os PDFs da pasta; o que não é comprovante é ignorado
    pid = _novo(cliente)
    pag = cliente.get(f'/p/{pid}/cnpjs').text
    for texto in ('Importar da pasta Orça.AI', 'na pasta "Orça.AI", em Documentos', 'Abrir a pasta Orça.AI', str(docs / 'Orça.AI')):
        assert texto in pag, texto
    assert 'Downloads' not in re.sub(r'<script.*?</script>', '', pag, flags=re.S)                # nenhum texto da tela fala mais em Downloads
    aberta = []
    monkeypatch.setattr('os.startfile', aberta.append, raising=False)
    r = cliente.post('/pasta-comprovantes/abrir', headers={'referer': f'http://x/p/{pid}/cnpjs'}, follow_redirects=False)
    assert r.status_code == 303 and r.headers['location'] == f'/p/{pid}/cnpjs' and aberta == [str(docs / 'Orça.AI')]
    r = cliente.post(f'/p/{pid}/cnpjs/importar', follow_redirects=False)
    assert 'Or%C3%A7a.AI' in r.headers['location'] and 'Downloads' not in r.headers['location']


# ---------------------------------------------------------------- faixa salarial pretendida do cargo
def _vaga_no_banco(V, cargo, empresa, cnpj, sal):
    from test_vagas import _pdf_vaga
    v = dict(titulo=cargo, empresa=empresa, url=f'https://vagas.exemplo/{V._slug(empresa)}', plataforma='InfoJobs', faixa_min=sal, faixa_max=sal, unidade='MONTH',
             cidade='São Paulo', uf='SP')
    V.guardar_no_banco(cargo, v, dict(status='🟢', cnpj=cnpj, razao_social=empresa.upper(), motivo='teste'),
                       _pdf_vaga('Salário R$ ' + f'{sal / 100:,.2f}'.replace(',', 'X').replace('.', ',').replace('X', '.')))
    return v['url']


def test_faixa_pretendida_escolhe_as_vagas_e_ajusta_as_horas(cliente):
    """Decisão da OSC (05/10/2026): o cargo tem uma faixa (quanto se quer pagar por mês); valem as vagas de MENOR salário cuja média chega
    nela, e as horas inteiras são ajustadas para o valor do plano ficar o mais perto possível da faixa."""
    from orcamento import db, vagas as V
    from orcamento.calculo import verificar, media_rh, horas_pela_faixa
    from orcamento.modelo import RubricaRH, PesquisaSalarial, Config
    from orcamento.otimizador import otimizar
    # as horas mais próximas da faixa: Assistente Social (30 h por semana = 150 h no mês), média 3.166,67 → 21,11 por hora
    r = RubricaRH(item=1, cargo='Assistente Social', horas_mes=40, meses=10, faixa_pretendida=200000,
                  pesquisas=[PesquisaSalarial(nome=n, cnpj=c, valor=v) for n, c, v in (('A', '1', 200000), ('B', '2', 350000), ('C', '3', 400000))])
    cfg = Config(horas_max_mes=220)                                                         # sem o limite de horas do projeto: só a jornada legal
    assert media_rh(r) == 316667 and horas_pela_faixa(r, cfg) == (95, 200545)             # 95 h = 2.005,45 (94 h daria 1.984,34: mais longe)
    assert horas_pela_faixa(r, Config()) == (90, 189990)                                   # com o limite padrão do projeto: no máximo 90 h por mês
    r.faixa_pretendida = 500000                                                             # acima da média: nem a jornada inteira alcança
    assert horas_pela_faixa(r, cfg) == (150, 316650) and horas_pela_faixa(r, Config()) == (90, 189990)
    r.faixa_pretendida = None
    assert horas_pela_faixa(r, cfg) is None
    # no sistema: cargo criado só com a faixa (sem horas); o banco tem 5 vagas
    pid = _novo(cliente)
    cliente.post(f'/p/{pid}/rubrica', data={'tipo': 'rh', 'nome': 'Psicólogo', 'faixa_pretendida': '2.000,00', 'meses': '10'})
    r = db.carregar(pid)[0].rubricas[0]
    assert r.faixa_pretendida == 200000 and r.horas_mes == 90                               # até haver vagas, o máximo de horas do projeto (padrão: 90 h)
    p0 = db.carregar(pid)[0]
    p0.config.horas_max_mes, p0.rubricas[0].horas_mes = 220, 220                            # (daqui em diante, sem o limite de 90 h: o caso é o da jornada inteira)
    db.salvar(pid, p0)
    for empresa, cnpj, sal in (('Alfa', '11111111000111', 150000), ('Beta', '22222222000122', 160000), ('Gama', '33333333000133', 170000),
                               ('Delta', '44444444000144', 210000), ('Epsilon', '55555555000155', 260000)):
        _vaga_no_banco(V, 'Psicólogo', empresa, cnpj, sal)
    assert [v['empresa'] for v in V.tres_do_titulo('Psicólogo')] == ['Alfa', 'Beta', 'Gama']              # sem faixa: as 3 de menor salário (média 1.600)
    tres = V.tres_na_faixa('Psicólogo', (), 200000)
    assert [v['empresa'] for v in tres] == ['Alfa', 'Delta', 'Epsilon'] and V.alcanca_a_faixa(tres, 200000)   # a menor média que chega em 2.000: 2.066,67
    assert [v['empresa'] for v in V.tres_na_faixa('Psicólogo', (), 300000)] == ['Alfa', 'Beta', 'Gama']    # nenhuma combinação chega em 3.000: as mais baratas
    pag = cliente.get(f'/p/{pid}/rh/1').text
    assert 'name="faixa_pretendida"' in pag and 'value="2.000,00"' in pag and 'de menor salário que chegam na faixa' in pag
    cliente.post(f'/p/{pid}/rh/1/vagas/banco')
    p = db.carregar(pid)[0]
    r = p.rubricas[0]
    assert [q.nome for q in r.pesquisas] == ['ALFA', 'DELTA', 'EPSILON'] and media_rh(r) == 206667
    assert (r.horas_mes, r.valor_mensal_plano) == (213, 200007)                             # 206.667 ÷ 220 = 9,39 por hora × 213 h = 2.000,07
    assert not [a for a in verificar(p) if a.regra == 'S09']
    pag = cliente.get(f'/p/{pid}/rh/1').text
    assert 'Faixa pretendida' in pag and '213 h' in pag and 'faixa pretendida R$ 2.000,00' in cliente.get(f'/p/{pid}').text
    # salvar a tela sem mexer não muda as horas; mudar a faixa recalcula
    from test_interface import _campos
    dados = _campos(pag)
    cliente.post(f'/p/{pid}/rh/1', data=dict(dados, horas_mes='200'))                       # horas mudadas à mão, faixa igual: ficam as da pessoa
    assert db.carregar(pid)[0].rubricas[0].horas_mes == 200
    cliente.post(f'/p/{pid}/rh/1', data=dict(dados, faixa_pretendida='1.000,00'))
    r = db.carregar(pid)[0].rubricas[0]
    assert (r.faixa_pretendida, r.horas_mes, r.valor_mensal_plano) == (100000, 106, 99534)                # 106 h = 995,34 (107 h = 1.004,73: mais longe)
    # faixa acima da média: jornada inteira e aviso na verificação
    cliente.post(f'/p/{pid}/rh/1', data=dict(dados, faixa_pretendida='5.000,00'))
    p = db.carregar(pid)[0]
    assert p.rubricas[0].horas_mes == 220 and [a.gravidade for a in verificar(p) if a.regra == 'S09'] == ['atencao']
    # "Fechar no teto" fica o mais perto possível da faixa
    cliente.post(f'/p/{pid}/rh/1', data=dict(dados, faixa_pretendida='2.000,00'))
    p = db.carregar(pid)[0]
    p.teto = 939 * 210 * 10                                                                  # só fecha com 210 h (3 a menos que as da faixa)
    res = otimizar(p)
    assert res['status'] == 'OK' and res['projeto'].rubricas[0].horas_mes == 210


def test_fechar_no_teto_reparte_por_igual_quando_as_faixas_nao_cabem(cliente):
    """Caso real de 05/10/2026: as faixas somavam mais que o teto e o fechamento tirou 20 h de um cargo só (o de 3 profissionais, em que cada hora
    pesa mais no total): valor 43% abaixo da faixa, com os outros cargos intactos. Agora a diferença é repartida em proporção da faixa de cada um."""
    from orcamento import db
    from orcamento.calculo import verificar
    from orcamento.modelo import RubricaRH, PesquisaSalarial
    from orcamento.otimizador import otimizar
    pid = _novo(cliente)
    p = db.carregar(pid)[0]
    pesquisas = lambda: [PesquisaSalarial(nome=n, cnpj=c, valor=v, faixa_min=v) for n, c, v in (('A', '11.222.333/0001-81', 230000), ('B', '11.444.777/0001-61', 300000),
                                                                                              ('C', '45.997.418/0001-53', 300000))]
    # dois cargos de 44 h por semana, mesma média (2.766,67 → 12,58 por hora), mesma faixa de 1.200 (95 h); um deles com 3 profissionais
    p.rubricas = [RubricaRH(item=1, cargo='Psicólogo', quantidade=3, horas_mes=95, meses=10, faixa_pretendida=120000, valor_mensal_plano=119510, pesquisas=pesquisas()),
                  RubricaRH(item=2, cargo='Designer Gráfico', quantidade=1, horas_mes=95, meses=10, faixa_pretendida=120000, valor_mensal_plano=119510, pesquisas=pesquisas())]
    p.teto = 1258 * (30 * 85 + 10 * 85)                       # só cabem 85 h em média: 10 h a menos que as da faixa (várias combinações fecham no centavo)
    p.config.horas_max_mes = 220                              # (o limite de 90 h do projeto não entra neste caso: a faixa pede 95 h)
    db.salvar(pid, p)
    res = otimizar(p)
    assert res['status'] == 'OK' and [r.horas_mes for r in res['projeto'].rubricas] == [85, 85]      # antes: 82 h e 94 h (quase tudo no cargo de 3 profissionais)
    assert [x['pct'] for x in res['faixas']] == [-10.9, -10.9] and res['sobra_das_faixas']['excesso'] == 1258 * 10 * 40
    assert [a.gravidade for a in verificar(res['projeto']) if a.regra == 'S09'] == ['info', 'info']  # a verificação diz que o valor está longe da faixa
    pag = cliente.post(f'/p/{pid}/otimizar', data={'ajustar_horas': '1'}).text
    assert 'As faixas não cabem no teto' in pag and 'Diferença para a faixa' in pag and '85 h' in pag and '-10,9%' in pag
    # quando as faixas cabem, cada cargo fica na hora mais próxima da faixa
    p.teto = 1258 * 95 * 40
    assert [r.horas_mes for r in otimizar(p)['projeto'].rubricas] == [95, 95]


def test_alt_clique_nao_baixa_a_pagina(cliente):
    """Caso real de 05/10/2026: no Chrome, Alt + clique num link BAIXA a página ("6.htm", "7.htm" em Downloads). O sistema trata o Alt + clique
    como clique comum (links e botões de enviar); só os links que abrem em outra aba ficam como o navegador faz. O sistema nunca manda baixar uma tela."""
    pid = _novo(cliente)
    r = cliente.get(f'/p/{pid}')
    assert 'content-disposition' not in {k.lower() for k in r.headers} and '<a download' not in r.text and ' download>' not in r.text
    js = r.text
    assert "if (!e.altKey || e.ctrlKey || e.shiftKey || e.metaKey || e.button !== 0) { return; }" in js
    assert "a.target === '_blank' || a.hasAttribute('download') || a.origin !== location.origin" in js and 'location.assign(a.href)' in js and 'reenviar(b.form, b)' in js


def test_horas_por_mes_nunca_passam_do_limite_do_projeto(cliente):
    """Decisão da OSC (06/10/2026): no cálculo automático das horas (faixa pretendida e "Fechar no teto"), as horas mensais de um cargo não
    passam de 90, e nunca da jornada definida em lei para o cargo. Horas digitadas acima do limite aparecem como erro na verificação."""
    from orcamento import db
    from orcamento.calculo import verificar, horas_pela_faixa
    from orcamento.modelo import RubricaRH, PesquisaSalarial, Config
    from orcamento.otimizador import otimizar
    from orcamento.regras import horas_maximas
    assert Config().horas_max_mes == 90 and horas_maximas('Psicólogo', Config()) == 90 and horas_maximas('Assistente Social', Config(horas_max_mes=200)) == 150
    assert horas_maximas('Auxiliar Administrativo', Config(horas_max_mes=300)) == 220 and horas_maximas('Técnico em Radiologia', Config(horas_max_mes=130)) == 120
    pesquisas = lambda: [PesquisaSalarial(nome=n, cnpj=c, valor=v, faixa_min=v) for n, c, v in (('A', '11.222.333/0001-81', 130000), ('B', '11.444.777/0001-61', 150000),
                                                                                              ('C', '45.997.418/0001-53', 170000))]
    # média 1.500 → 6,82 por hora (44 h semanais = 220 h): a faixa de 1.000 pediria 147 h; com o limite, 90 h (613,80)
    r = RubricaRH(item=1, cargo='Auxiliar Administrativo', horas_mes=40, meses=10, faixa_pretendida=100000, pesquisas=pesquisas())
    assert horas_pela_faixa(r, Config(horas_max_mes=220)) == (147, 100254) and horas_pela_faixa(r, Config()) == (90, 61380)
    pid = _novo(cliente)
    p = db.carregar(pid)[0]
    p.rubricas = [RubricaRH(item=1, cargo='Auxiliar Administrativo', horas_mes=129, meses=10, valor_mensal_plano=87978, pesquisas=pesquisas())]
    p.teto = 682 * 80 * 10                                             # o teto cabe em 80 h
    db.salvar(pid, p)
    s12 = [a for a in verificar(p) if a.regra == 'S12']
    assert len(s12) == 1 and s12[0].gravidade == 'erro' and '129 horas por mês: o máximo para este cargo é 90 h' in s12[0].mensagem
    res = otimizar(p)
    assert res['status'] == 'OK' and res['projeto'].rubricas[0].horas_mes == 80 and not [a for a in verificar(res['projeto']) if a.regra == 'S12']
    p.teto = 682 * 150 * 10                                            # um teto que só fecharia com 150 h: o "Fechar no teto" não passa das 90 h
    assert otimizar(p)['status'] != 'OK'
    p.teto = 682 * 80 * 10
    # o limite é do projeto: dá para mudar na configuração (e nunca passa da jornada legal do cargo)
    pag = cliente.get(f'/p/{pid}').text
    assert 'name="horas_max_mes"' in pag and 'Máximo de horas por mês de um cargo' in pag
    p.config.horas_max_mes = 130
    db.salvar(pid, p)
    assert not [a for a in verificar(db.carregar(pid)[0]) if a.regra == 'S12']


def test_com_o_limite_de_horas_as_vagas_sao_escolhidas_pela_media_que_chega_na_faixa(cliente):
    """Pesquisa completa de 09/10/2026: "Auxiliar de Serviços Gerais", faixa de R$ 1.000 — as 3 vagas de menor salário tinham média de R$ 1.673,67 e,
    com o máximo de 90 h, o valor ficou em R$ 684,90. Com o limite de horas, a média que as vagas precisam ter é faixa × horas do mês ÷ máximo de
    horas (R$ 2.444,45 para 220 h e 90 h): é com ela que as vagas são procuradas e escolhidas."""
    from orcamento import db, vagas as V
    from orcamento.calculo import media_para_a_faixa, media_rh, verificar
    from orcamento.modelo import RubricaRH, Config
    r = RubricaRH(item=1, cargo='Auxiliar Administrativo', horas_mes=40, meses=10, faixa_pretendida=100000)
    assert media_para_a_faixa(r, Config()) == 244445 and media_para_a_faixa(r, Config(horas_max_mes=220)) == 100000      # sem o limite, a média é a própria faixa
    assert media_para_a_faixa(RubricaRH(item=1, cargo='Assistente Social', horas_mes=40, meses=10, faixa_pretendida=100000), Config()) == 166667   # 150 h no mês
    assert media_para_a_faixa(RubricaRH(item=1, cargo='Psicólogo', horas_mes=40, meses=10), Config()) is None
    pid = _novo(cliente)
    cliente.post(f'/p/{pid}/rubrica', data={'tipo': 'rh', 'nome': 'Auxiliar Administrativo', 'faixa_pretendida': '1.000,00', 'meses': '10'})
    for empresa, cnpj, sal in (('Alfa', '11111111000111', 150000), ('Beta', '22222222000122', 160000), ('Gama', '33333333000133', 170000),
                               ('Delta', '44444444000144', 260000), ('Epsilon', '55555555000155', 270000), ('Zeta', '66666666000166', 280000)):
        _vaga_no_banco(V, 'Auxiliar Administrativo', empresa, cnpj, sal)
    assert [v['empresa'] for v in V.tres_do_titulo('Auxiliar Administrativo', 100000)] == ['Alfa', 'Beta', 'Gama']          # pela faixa sozinha: média de 1.600
    assert [v['empresa'] for v in V.tres_do_titulo('Auxiliar Administrativo', 244445)] == ['Delta', 'Epsilon', 'Zeta']       # pela média necessária: 2.700
    pag = cliente.get(f'/p/{pid}/rh/1').text
    assert 'a média das 3 precisa ser de pelo menos R$ 2.444,45' in pag
    cliente.post(f'/p/{pid}/rh/1/vagas/banco')
    p = db.carregar(pid)[0]
    r = p.rubricas[0]
    assert [q.nome for q in r.pesquisas] == ['DELTA', 'EPSILON', 'ZETA'] and media_rh(r) == 270000
    assert r.horas_mes <= 90 and abs(r.valor_mensal_plano - 100000) <= 1227                                              # chega na faixa (a menos de uma hora de diferença), dentro das 90 h
    assert not [a for a in verificar(p) if a.regra in ('S09', 'S12') and a.gravidade != 'info']


def test_aviso_diz_quando_um_titulo_similar_chega_na_faixa(cliente, monkeypatch):
    """Pesquisa completa de 09/10/2026: "Auxiliar Administrativo", faixa de R$ 1.200 — nenhuma combinação de vagas do título chegava na média
    necessária (R$ 2.933,34), mas o título similar "Assistente administrativo" tinha 3 vagas que chegavam, e o aviso não dizia. As vagas de
    títulos diferentes não se misturam (decisão da OSC): o sistema mantém as do título do cargo e aponta o similar como opção."""
    import asyncio
    from orcamento import db, servico, vagas as V
    pid = _novo(cliente)
    cliente.post(f'/p/{pid}/rubrica', data={'tipo': 'rh', 'nome': 'Auxiliar Administrativo', 'faixa_pretendida': '1.200,00', 'meses': '10'})
    for cargo, empresa, cnpj, sal in (('Auxiliar Administrativo', 'Alfa', '11111111000111', 120000), ('Auxiliar Administrativo', 'Beta', '22222222000122', 141200),
                                      ('Auxiliar Administrativo', 'Gama', '33333333000133', 151800), ('Assistente administrativo', 'Delta', '44444444000144', 300000),
                                      ('Assistente administrativo', 'Epsilon', '55555555000155', 308800), ('Assistente administrativo', 'Zeta', '66666666000166', 310000)):
        _vaga_no_banco(V, cargo, empresa, cnpj, sal)

    async def sem_busca(cargo, ctx=None, alvo=3, parar=None, faixa=None):
        return dict(cargo=cargo, lidas=0, cnpj_consultados=0, cnpj_verdes=0, no_banco=len(V.tres_do_banco(cargo)), vagas=V.tres_do_banco(cargo))
    monkeypatch.setattr(V, 'coletar', sem_busca)

    class Ctx:
        id, avisos = None, []
        def aviso(self, m): self.avisos.append(m)
        def etapa(self, m): pass
        def progresso(self, *a): pass
    ctx = Ctx()
    asyncio.run(servico.vagas_do_cargo(pid, 1, ctx))
    r = db.carregar(pid)[0].rubricas[0]
    assert [q.nome for q in r.pesquisas] == ['ALFA', 'BETA', 'GAMA'] and r.horas_mes == 90          # ficam as do título do cargo (as mais baratas), no máximo de horas
    aviso = next(a for a in ctx.avisos if 'não chegam na faixa pretendida' in a)
    assert 'a média precisaria ser de R$ 2.933,34' in aviso
    assert 'Título similar com 3 vagas que CHEGAM na faixa: Assistente administrativo (salários de R$ 3.000,00 a R$ 3.100,00)' in aviso and 'Títulos com vagas' in aviso
