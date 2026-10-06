"""Órgãos com regras dinâmicas (pedido de 05/10/2026): o sistema serve a qualquer secretaria. Cada órgão é um cadastro feito pela tela, com as
regras do sistema que valem para ele (ligada/desligada, peso), as regras próprias (tipos que o sistema sabe conferir, preenchidos pela pessoa,
ou texto conferido por uma pessoa ou pela IA) e a forma de fazer o orçamento. O projeto segue as regras do órgão dele."""
import json
import re

import pytest
from fastapi.testclient import TestClient


@pytest.fixture
def cliente(tmp_path, monkeypatch):
    monkeypatch.setenv('ORCAMENTO_DADOS', str(tmp_path))
    import importlib, orcamento.db as dbm
    monkeypatch.setattr(dbm, 'PASTA_DADOS', str(tmp_path))
    monkeypatch.setattr(dbm, 'PASTA_LOCAL', str(tmp_path / 'local'))
    import app as appmod
    importlib.reload(appmod)
    return TestClient(appmod.app)


def _orgao(cliente, nome='Secretaria Municipal de Assistência Social', sigla='SMAS', **mais):
    r = cliente.post('/orgaos', data=dict(dict(nome=nome, sigla=sigla, esfera='municipal', uf='sp'), **mais), follow_redirects=False)
    return int(re.search(r'/orgaos/(\d+)', r.headers['location']).group(1))


def _projeto(cliente, oid=None):
    dados = {'nome': 'T', 'teto': '10.000,00', 'cep': '01001-000'}
    if oid:
        dados['orgao_id'] = str(oid)
    return int(cliente.post('/projetos', data=dados, follow_redirects=False).headers['location'].rsplit('/', 1)[-1])


def _plano(pid):
    """Um cargo (3 vagas, 2 de SP e 1 do RJ) e uma rubrica com dois itens (um deles, um notebook)."""
    from orcamento import db
    from orcamento.modelo import RubricaRH, RubricaMaterial, PesquisaSalarial, Subitem, Fonte, Evidencia
    p = db.carregar(pid)[0]
    ev = lambda: Evidencia(arquivo='x.pdf', origem='pdf', sha256='a' * 64)
    p.rubricas = [
        RubricaRH(item=1, cargo='Psicólogo', horas_mes=100, meses=10, valor_mensal_plano=100000,
                  pesquisas=[PesquisaSalarial(nome=n, cnpj=c, valor=v, faixa_min=v, data_pesquisa='2026-10-01', evidencia=ev())
                             for n, c, v in (('ALFA', '11.222.333/0001-81', 220000), ('BETA', '11.444.777/0001-61', 220000), ('GAMA', '45.997.418/0001-53', 220000))]),
        RubricaMaterial(item=2, descricao='Equipamentos e alimentação', meses=14, subitens=[
            Subitem(descricao='Notebook', especificacao='15 polegadas', qtd=1, precos=[300000, 320000, 340000], valor_plano=300000,
                    fontes=[Fonte(nome=f'LOJA {k}', cnpj=c, data_pesquisa='2026-10-01', evidencia=ev()) for k, c in enumerate(('11.222.333/0001-81', '11.444.777/0001-61', '45.997.418/0001-53'))]),
            Subitem(descricao='Café', especificacao='500g', qtd=2, precos=[1000, 1200, 2000], valor_plano=1000,
                    fontes=[Fonte(nome=f'LOJA {k}', cnpj=c, data_pesquisa='2026-10-01', evidencia=ev()) for k, c in enumerate(('11.222.333/0001-81', '11.444.777/0001-61', '45.997.418/0001-53'))])])]
    db.salvar(pid, p)
    return p


def test_orgao_padrao_e_projeto_ligado_ao_orgao(cliente):
    from orcamento import db, orgaos
    from orcamento.calculo import verificar
    pag = cliente.get('/orgaos').text
    assert 'Secretaria da Justiça e Cidadania do Estado de São Paulo' in pag and 'Padrão do sistema' in pag and 'Adicionar órgão' in pag
    padrao = orgaos.padrao()
    assert orgaos.ler(padrao).sigla == 'SEJC-SP' and orgaos.padrao() == padrao                       # criado uma vez só
    pid = _projeto(cliente)                                                                         # sem escolher: o órgão padrão
    assert db.carregar(pid)[0].orgao_id == padrao
    oid = _orgao(cliente)
    o = orgaos.ler(oid)
    assert (o.nome, o.sigla, o.esfera, o.uf) == ('Secretaria Municipal de Assistência Social', 'SMAS', 'municipal', 'SP') and not o.regras and not o.proprias
    pid2 = _projeto(cliente, oid)
    p = db.carregar(pid2)[0]
    assert p.orgao_id == oid and p.orgao == o.nome
    pag = cliente.get(f'/p/{pid2}').text
    assert 'com as regras de <b>Secretaria Municipal de Assistência Social</b>' in pag and '· SMAS ·' in pag and f'href="/orgaos/{oid}"' in pag
    # projeto antigo, sem órgão: nada muda na verificação e, ao abrir, ele passa a pertencer ao órgão padrão
    p.orgao_id = None
    antes = [(a.regra, a.gravidade, a.item) for a in verificar(p)]
    db.salvar(pid2, p)
    cliente.get(f'/p/{pid2}')
    p = db.carregar(pid2)[0]
    assert p.orgao_id == padrao and [(a.regra, a.gravidade, a.item) for a in verificar(p)] == antes
    # trocar o órgão pela configuração do projeto
    from test_interface import _campos
    cliente.post(f'/p/{pid2}/config', data=dict(_campos(cliente.get(f'/p/{pid2}').text), orgao_id=str(oid)))
    assert db.carregar(pid2)[0].orgao_id == oid


def test_regras_do_sistema_ligadas_desligadas_e_com_outro_peso(cliente):
    from orcamento import db, orgaos
    from orcamento.calculo import verificar
    oid = _orgao(cliente)
    pid = _projeto(cliente, oid)
    p = _plano(pid)
    p.rubricas[0].pesquisas[1].cnpj = p.rubricas[0].pesquisas[0].cnpj          # a mesma empresa em duas pesquisas (R09) ...
    p.rubricas[0].valor_mensal_plano = 90000                                    # ... e o plano fora do teto (R05)
    db.salvar(pid, p)
    regras = lambda: {(a.regra, a.gravidade) for a in verificar(db.carregar(pid)[0])}
    assert ('R09', 'erro') in regras() and any(r == 'R05' for r, _ in regras())
    pag = cliente.get(f'/orgaos/{oid}').text
    assert 'name="ativa_R09"' in pag and 'name="grav_R09"' in pag and 'name="ativa_R05"' not in pag     # a do teto vale para qualquer órgão
    from test_interface import _campos
    dados = _campos(pag)
    cliente.post(f'/orgaos/{oid}', data=dict(dados, grav_R09='atencao'))                               # abrandar: vira ponto para revisar
    assert ('R09', 'atencao') in regras() and orgaos.ler(oid).regras['R09'].gravidade == 'atencao'
    dados.pop('ativa_R09'); dados['ativa_R05'] = ''                                                     # desligar a R09; tentar desligar a R05 não tem efeito
    cliente.post(f'/orgaos/{oid}', data=dados)
    assert not any(r == 'R09' for r, _ in regras()) and any(r == 'R05' for r, _ in regras())
    assert set(orgaos.ler(oid).regras) == {'R09'} and not orgaos.ler(oid).regras['R09'].ativa          # só o que difere do padrão fica guardado
    # o órgão padrão continua com tudo ligado: o mesmo plano, num projeto dele, ainda tem a R09 como pendência
    pid2 = _projeto(cliente)
    p2 = db.carregar(pid2)[0]; p2.rubricas = db.carregar(pid)[0].rubricas; db.salvar(pid2, p2)
    assert ('R09', 'erro') in {(a.regra, a.gravidade) for a in verificar(db.carregar(pid2)[0])}
    # a origem da regra, nas telas, é o órgão do projeto
    assert orgaos.texto_da_regra('R02', orgaos.ler(oid))[1] == 'regra de SMAS' and orgaos.texto_da_regra('R02', orgaos.padrao_embutido())[1] == 'regra da SEJC'


def test_regras_proprias_de_cada_tipo(cliente, monkeypatch):
    from orcamento import db, orgaos, regras_dinamicas as RD
    from orcamento.calculo import verificar
    monkeypatch.setattr(RD, '_uf_do_cnpj', lambda c: {'11.222.333/0001-81': 'SP', '11.444.777/0001-61': 'RJ'}.get(c))   # o 3º CNPJ não está na base
    oid = _orgao(cliente)
    pid = _projeto(cliente, oid)
    _plano(pid)
    nova = lambda **d: cliente.post(f'/orgaos/{oid}/regra', data=dict(dict(ativa='1'), **d), follow_redirects=False).headers['location']
    assert 'P01' in nova(tipo='uf_empresas', titulo='Só empresas do estado', gravidade='erro', p_ufs='sp', p_aplica='rh')
    assert 'P02' in nova(tipo='palavras_proibidas', titulo='Material permanente não pode', gravidade='erro', p_palavras='notebook, impressoras, ar-condicionado', p_alvo='itens')
    assert 'P03' in nova(tipo='percentual_maximo', titulo='Mão de obra até 15% do total', gravidade='atencao', p_grupo='rh', p_pct='15')
    assert 'P04' in nova(tipo='valor_maximo', titulo='Item até R$ 2.000', gravidade='atencao', p_alvo='item', p_valor='2.000,00', p_texto='')
    assert 'P05' in nova(tipo='duracao_maxima', titulo='Até 12 meses', gravidade='erro', p_meses='12')
    assert 'P06' in nova(tipo='lembrete', titulo='Assinatura do responsável técnico', gravidade='atencao', p_texto='O plano de trabalho precisa da assinatura do responsável técnico.')
    assert 'P07' in nova(tipo='ia', titulo='Coffee break até R$ 20 por pessoa', gravidade='atencao', p_texto='Não são permitidas despesas com coffee break acima de R$ 20,00 por pessoa.')
    assert [r.codigo for r in orgaos.ler(oid).proprias] == [f'P0{n}' for n in range(1, 8)] and orgaos.ler(oid).proprias[3].parametros['valor'] == 200000
    assert 'gravado' in nova(tipo='uf_empresas', titulo='x', gravidade='erro', p_ufs='XX') and len(orgaos.ler(oid).proprias) == 7     # estado que não existe: não grava
    al = [(a.regra, a.gravidade, a.item, a.mensagem) for a in verificar(db.carregar(pid)[0]) if a.regra.startswith('P')]
    por = lambda cod: [(g, onde, msg) for r, g, onde, msg in al if r == cod]
    assert [x[:2] for x in por('P01')] == [('erro', 'Item 1 – Psicólogo'), ('atencao', 'Item 1 – Psicólogo')]          # BETA é do RJ; de GAMA não se sabe o estado
    assert 'BETA' in por('P01')[0][2] and 'é de RJ' in por('P01')[0][2] and 'GAMA' in por('P01')[1][2]
    assert [x[:2] for x in por('P02')] == [('erro', 'Item 2 – Equipamentos e alimentação / Notebook')] and '"notebook" não pode constar' in por('P02')[0][2]
    assert por('P03')[0][:2] == ('atencao', 'Plano') and '% do total' in por('P03')[0][2]                                # mão de obra: 10.000 de 52.280
    assert [x[1] for x in por('P04')] == ['Item 2 – Equipamentos e alimentação / Notebook'] and 'R$ 3.000,00' in por('P04')[0][2]
    assert [x[:2] for x in por('P05')] == [('erro', 'Item 2 – Equipamentos e alimentação')] and '14 meses' in por('P05')[0][2]
    assert por('P06') == [('atencao', 'Plano', 'O plano de trabalho precisa da assinatura do responsável técnico. (conferência por uma pessoa)')]
    assert por('P07')[0][0] == 'info' and 'ainda não foi conferida pela IA' in por('P07')[0][2]                           # a tela nunca consulta a IA
    # telas: as regras aparecem no órgão e na verificação do projeto, com o nome que a pessoa deu
    pag = cliente.get(f'/orgaos/{oid}').text
    assert 'id="regra-P02"' in pag and '3 palavra(s) proibida(s)' in pag and 'empresas de SP (mão de obra)' in pag and 'cada item até R$ 2.000,00' in pag
    proj = cliente.get(f'/p/{pid}').text
    assert 'Material permanente não pode' in proj and 'regra própria de SMAS' in proj and 'Conferir regras com a IA (1)' in proj
    # a IA responde: o resultado fica guardado para este plano e vira ponto para revisar; se o plano muda, volta a "não conferida"
    from orcamento import ia
    monkeypatch.setattr(ia, 'conferir_regra', lambda regra, plano, pid_=None: dict(violacoes=[dict(n=2, motivo='café acima do limite por pessoa')]))
    p = db.carregar(pid)[0]
    assert RD.conferir_com_a_ia(p, orgaos.ler(oid), pid) == (1, 0) and RD.conferir_com_a_ia(p, orgaos.ler(oid), pid) == (0, 0)
    assert [(a.gravidade, a.item) for a in verificar(p) if a.regra == 'P07'] == [('atencao', 'Item 2 – Equipamentos e alimentação')]
    p.rubricas[1].subitens[1].qtd = 3
    assert [a.gravidade for a in verificar(p) if a.regra == 'P07'] == ['info']
    # editar, desligar e excluir
    cliente.post(f'/orgaos/{oid}/regra', data=dict(codigo='P05', tipo='duracao_maxima', titulo='Até 24 meses', gravidade='erro', p_meses='24', ativa='1'))
    assert not [a for a in verificar(db.carregar(pid)[0]) if a.regra == 'P05']
    cliente.post(f'/orgaos/{oid}/regra', data=dict(codigo='P02', tipo='palavras_proibidas', titulo='Material permanente não pode', gravidade='erro', p_palavras='notebook', p_alvo='itens'))
    assert not orgaos.ler(oid).proprias[1].ativa and not [a for a in verificar(db.carregar(pid)[0]) if a.regra == 'P02']      # sem "ativa": desligada
    cliente.post(f'/orgaos/{oid}/regra/P06/excluir')
    assert 'P06' not in [r.codigo for r in orgaos.ler(oid).proprias] and orgaos.ler(oid).proximo_codigo() == 'P06'
    # marcar o lembrete como revisado funciona como nos outros pontos (o órgão padrão não tem nada disso)
    assert not [a for a in verificar(db.carregar(_projeto(cliente))[0]) if a.regra.startswith('P')]


def test_ia_propoe_regras_a_partir_do_texto_e_a_pessoa_decide(cliente, monkeypatch):
    from orcamento import ia, orgaos
    oid = _orgao(cliente)
    texto = 'As cotações devem ser realizadas com fornecedores sediados no Estado de São Paulo. Não é permitida a aquisição de material permanente. Dispensa-se a comprovação da página da vaga.'
    proposta = dict(proprias=[
        dict(tipo='uf_empresas', titulo='Fornecedores de São Paulo', gravidade='erro', parametros=dict(ufs=['SP'], aplica='todos'), trecho='fornecedores sediados no Estado de São Paulo'),
        dict(tipo='palavras_proibidas', titulo='Sem material permanente', gravidade='erro', parametros=dict(palavras='computador, impressora', alvo='itens'), trecho='Não é permitida a aquisição de material permanente'),
        dict(tipo='tipo_que_nao_existe', titulo='x', parametros={}), dict(tipo='uf_empresas', titulo='estado inválido', parametros=dict(ufs=['ZZ'])),
        dict(tipo='valor_maximo', titulo='sem valor', parametros=dict(alvo='item'))],
        desligar=[dict(codigo='S05', trecho='Dispensa-se a comprovação da página da vaga'), dict(codigo='R05', trecho='x'), dict(codigo='XYZ', trecho='x')])
    monkeypatch.setattr(ia, 'disponivel', lambda: True)
    monkeypatch.setattr(ia, 'propor_regras', lambda t, tipos, catalogo, pid=None: proposta if 'São Paulo' in t else None)
    pag = cliente.post(f'/orgaos/{oid}/ia', data={'texto': texto}).text
    assert 'O que a IA propõe a partir do texto' in pag and 'Fornecedores de São Paulo' in pag and 'Sem material permanente' in pag and 'Desligar <b>S05</b>' in pag
    assert 'fornecedores sediados no Estado de São Paulo' in pag                                           # o trecho de onde a regra saiu
    for fora in ('tipo_que_nao_existe', 'estado inválido', 'sem valor', 'Desligar <b>R05</b>', 'XYZ'):     # o que o sistema não consegue usar é descartado
        assert fora not in pag, fora
    assert not orgaos.ler(oid).proprias and not orgaos.ler(oid).regras                                      # nada gravado: é só uma proposta
    dados = re.search(r'name="dados" value="([^"]*)"', pag).group(1).replace('&#34;', '"').replace('&amp;', '&')
    assert len(json.loads(dados)['proprias']) == 2
    r = cliente.post(f'/orgaos/{oid}/ia/aplicar', data={'dados': dados, 'propria': ['0'], 'desligar': ['S05']}, follow_redirects=False)   # a pessoa aceita uma das duas
    assert r.status_code == 303
    o = orgaos.ler(oid)
    assert [(x.codigo, x.tipo, x.titulo, x.gravidade, x.parametros['ufs']) for x in o.proprias] == [('P01', 'uf_empresas', 'Fornecedores de São Paulo', 'erro', ['SP'])]
    assert set(o.regras) == {'S05'} and not o.regras['S05'].ativa
    assert 'Não foi possível consultar a IA' in cliente.post(f'/orgaos/{oid}/ia', data={'texto': 'um texto qualquer do órgão, comprido o bastante para ser lido'}).text
    assert 'cole um trecho maior' in cliente.post(f'/orgaos/{oid}/ia', data={'texto': 'curto'}).text


def test_remover_restaurar_e_copiar_orgao(cliente):
    from orcamento import orgaos
    oid = _orgao(cliente)
    cliente.post(f'/orgaos/{oid}/regra', data=dict(tipo='duracao_maxima', titulo='Até 12 meses', gravidade='erro', p_meses='12', ativa='1'))
    copia = _orgao(cliente, nome='Fundo Municipal', sigla='FMDCA', copiar=str(oid))                        # órgão novo a partir das regras de outro
    assert [r.titulo for r in orgaos.ler(copia).proprias] == ['Até 12 meses'] and orgaos.ler(copia).nome == 'Fundo Municipal'
    pid = _projeto(cliente, oid)
    padrao = orgaos.padrao()
    assert 'não pode ser removido' in cliente.post(f'/orgaos/{padrao}/remover', follow_redirects=True).text and padrao in [x['id'] for x in orgaos.listar()]
    cliente.post(f'/orgaos/{oid}/remover')
    assert oid not in [x['id'] for x in orgaos.listar()] and oid in [x['id'] for x in orgaos.listar(removidos=True)]
    assert orgaos.ler(oid) is not None and 'Até 12 meses' in cliente.get(f'/orgaos/{oid}').text           # nada é apagado: o projeto dele continua com as regras
    assert 'com as regras de <b>Secretaria Municipal de Assistência Social</b>' in cliente.get(f'/p/{pid}').text
    assert 'Órgãos removidos (1)' in cliente.get('/orgaos').text
    cliente.post(f'/orgaos/{oid}/restaurar')
    assert oid in [x['id'] for x in orgaos.listar()]
    assert cliente.get('/orgaos/9999').status_code in (404, 500)


def test_planilha_traz_todas_as_abas_do_modelo_e_segue_o_orgao(cliente):
    """Pedido da OSC (05/10/2026): a planilha do pacote traz todas as abas da planilha de pré-cálculos — além do Plano e do Comparativo, o
    Cronograma físico-financeiro, as Etapas e Fases e o Cronograma de desembolso. Os cronogramas puxam os valores do Plano por fórmula."""
    import io
    from openpyxl import load_workbook
    from orcamento import db, orgaos
    from orcamento.calculo import periodo, duracao_do_projeto
    from test_interface import _campos
    pid = _projeto(cliente)
    p = _plano(pid)                                                            # cargo de 10 meses (R$ 1.000 por mês) e rubrica de 14 meses (R$ 3.020 por mês)
    assert duracao_do_projeto(p) == 14 and periodo(p, p.rubricas[0]) == (3, 12) and periodo(p, p.rubricas[1]) == (1, 14)   # o mais curto fica no meio do projeto
    abrir = lambda: (load_workbook(io.BytesIO(cliente.get(f'/p/{pid}/planilhas').content)), load_workbook(io.BytesIO(cliente.get(f'/p/{pid}/planilhas').content), data_only=True))
    wb, sem = abrir()
    assert wb.sheetnames == ['Plano de Aplicação', 'Cronograma fisico-financeiro', 'Etapa e Fases', 'Cronograma de desembolso', 'Comparativo de Preço']
    c, cv = wb['Cronograma fisico-financeiro'], sem['Cronograma fisico-financeiro']
    assert [c[x].value for x in ('A1', 'A2', 'B2', 'C2', 'P2', 'A3', 'A5', 'A6')] == ['CRONOGRAMA FÍSICO-FINANCEIRO', 'Item', 'Descrição', '1º MÊS', '14º MÊS', 1, 'TOTAL', 'TOTAL AGRUPADO']
    assert [c[x].value for x in ('B3', 'D3', 'E3', 'F3', 'N3', 'O3', 'C4', 'D4', 'P4')] == \
        ["='Plano de Aplicação'!D4", None, "='Plano de Aplicação'!E4", '=$E$3', '=$E$3', None, "='Plano de Aplicação'!E5", '=$C$4', '=$C$4']   # o valor mensal só nos meses da rubrica
    assert [c[x].value for x in ('C5', 'C6', 'D6', 'Q3', 'Q6', 'Q7')] == ['=SUM(C3:C4)', '=C5', '=C6+D5', '=SUM(C3:P3)', '=SUM(Q3:Q4)', '=10000.0-Q6']
    assert [cv[x].value for x in ('B3', 'E3', 'N3', 'C4', 'C5', 'E5', 'C6', 'D6', 'E6', 'P6', 'Q3', 'Q4', 'Q6', 'Q7')] == \
        ['Psicólogo (recibo ou MEI) - 100h mensal', 1000, 1000, 3020, 3020, 4020, 3020, 6040, 10060, 52280, 10000, 42280, 52280, -42280]
    assert (c['A2'].font.name, c['A2'].font.sz, c['A2'].font.b, c['A2'].fill.fgColor.rgb, c['B3'].font.name) == ('Verdana', 5, True, 'FFE7E6E6', 'Times New Roman')
    assert {str(m) for m in c.merged_cells.ranges} == {'A1:P1', 'A5:B5', 'A6:B6'} and c.print_area.endswith('$A$1:$P$6') and c.page_setup.orientation == 'landscape'
    assert c['C3'].number_format.startswith('_-"R$ "* #,##0.00') and c['D3'].border.left.style == 'medium'      # o mês vazio também tem a borda
    e, ev = wb['Etapa e Fases'], sem['Etapa e Fases']
    assert [e[x].value for x in ('A1', 'A2', 'B2', 'C2', 'D2', 'A3', 'C3', 'D3', 'C4', 'D4')] == \
        ['ETAPAS E FASES', 'Item', 'Etapa', 'Atividade', 'Prazo', 1, 'Recursos Humanos', '3º ao 12º mês', 'Recursos Materiais', '1º ao 14º mês']
    assert (ev['B3'].value, ev['B4'].value, e['A1'].font.name, e['A1'].font.sz, e['B3'].font.name) == ('Psicólogo (recibo ou MEI) - 100h mensal', 'Equipamentos e alimentação', 'Verdana', 11, 'Times New Roman')
    d, dv = wb['Cronograma de desembolso'], sem['Cronograma de desembolso']
    assert [d[x].value for x in ('A1', 'A2', 'N2', 'A4', 'B4')] == ['CRONOGRAMA DE DESEMBOLSO', '1º mês', '14º mês', "='Plano de Aplicação'!G8", None] and dv['A4'].value == 52280   # tudo no 1º mês
    assert {'A1:N1', 'A2:A3', 'N2:N3'} <= {str(m) for m in d.merged_cells.ranges}
    for ws in wb.worksheets:   # nenhuma fórmula sem o valor gravado (quem abre sem recalcular vê os números e os nomes)
        for linha in ws.iter_rows():
            for cel in linha:
                if isinstance(cel.value, str) and cel.value.startswith('='):
                    assert sem[ws.title][cel.coordinate].value is not None, (ws.title, cel.coordinate)
    # o mês de início é escolhido na tela do cargo e da rubrica (em branco, volta a ser calculado)
    pag = cliente.get(f'/p/{pid}/rh/1').text
    assert 'name="mes_inicio"' in pag and 'hoje, 3º ao 12º mês' in pag and 'name="mes_inicio"' in cliente.get(f'/p/{pid}/mat/2').text
    cliente.post(f'/p/{pid}/rh/1', data=dict(_campos(pag), mes_inicio='5'))
    p = db.carregar(pid)[0]
    assert p.rubricas[0].mes_inicio == 5 and periodo(p, p.rubricas[0]) == (5, 14)
    wb, sem = abrir()
    assert wb['Etapa e Fases']['D3'].value == '5º ao 14º mês' and wb['Cronograma fisico-financeiro']['G3'].value == "='Plano de Aplicação'!E4" and wb['Cronograma fisico-financeiro']['F3'].value is None
    cliente.post(f'/p/{pid}/rh/1', data=dict(_campos(cliente.get(f'/p/{pid}/rh/1').text), mes_inicio=''))
    assert db.carregar(pid)[0].rubricas[0].mes_inicio is None
    # o órgão do projeto: o nome na coluna do concedente, o desembolso mês a mês e, no modelo simples, só o Plano e o Comparativo
    oid = _orgao(cliente)
    o = orgaos.ler(oid)
    assert o.no_plano() == 'SMAS' and o.parametros.desembolso == 'unica' and o.modelo_planilha == 'sejc'
    cliente.post(f'/orgaos/{oid}', data=dict(_campos(cliente.get(f'/orgaos/{oid}').text), desembolso='mensal', concedente='Assistência Social'))
    assert orgaos.ler(oid).parametros.desembolso == 'mensal'
    cliente.post(f'/p/{pid}/config', data=dict(_campos(cliente.get(f'/p/{pid}').text), orgao_id=str(oid)))
    wb, sem = abrir()
    assert wb['Plano de Aplicação']['G2'].value == 'Concedente\n (Assistência Social)'
    assert [wb['Cronograma de desembolso'][x].value for x in ('A4', 'C4', 'N4')] == ["='Cronograma fisico-financeiro'!C5", "='Cronograma fisico-financeiro'!E5", "='Cronograma fisico-financeiro'!P5"]
    assert [sem['Cronograma de desembolso'][x].value for x in ('A4', 'C4', 'N4')] == [3020, 4020, 3020]
    cliente.post(f'/orgaos/{oid}', data=dict(_campos(cliente.get(f'/orgaos/{oid}').text), modelo_planilha='simples'))
    assert load_workbook(io.BytesIO(cliente.get(f'/p/{pid}/planilhas').content)).sheetnames == ['Plano de Aplicação', 'Comparativo de Preço']
    # o projeto do órgão padrão continua com "Concedente (SJC)" e as cinco abas
    pid2 = _projeto(cliente); _plano(pid2)
    wb2 = load_workbook(io.BytesIO(cliente.get(f'/p/{pid2}/planilhas').content))
    assert wb2['Plano de Aplicação']['G2'].value == 'Concedente\n (SJC)' and len(wb2.sheetnames) == 5


def test_altura_das_linhas_do_comparativo_acompanha_o_texto():
    """Na impressão pelo Excel, nomes compridos de empresa saíam cortados: a altura da linha era estimada por 15 caracteres por linha, mas o
    Excel quebra nas palavras e as maiúsculas ocupam mais. A estimativa agora confere com o que o Excel mostrou (medido em 05/10/2026, com nomes do mesmo formato)."""
    from orcamento.pacote import _linhas_quebradas as linhas
    for texto, no_excel in (('EMPRESA EXEMPLAR DE SANEAMENTO E SERVICOS LTDA - CNPJ: 11.222.333/0001-81', 6), ('ALFAXI TECNOLOGIA LTDA - CNPJ: 11.222.333/0001-81', 4),
                            ('CENTRO TERAPEUTICO INTEGRADO MULTIDISCIPLINAR EM SAUDE LTDA - CNPJ: 11.444.777/0001-61', 7), ('Auxiliar Administrativo (recibo ou MEI) - 101h mensal', 4),
                            ('ASSOCIACAO BENEFICENTE JARDINS DE SOL - ABEJAS - CNPJ: 11.222.333/0001-81', 5), ('', 1), ('Café', 1), ('A' * 40, 3)):
        assert linhas(texto) == no_excel, texto


def test_planilhas_em_pdf_saem_da_mesma_planilha(cliente):
    """05/10/2026: o Plano, os cronogramas e o Comparativo também em PDF. O PDF não é montado à parte: cada aba da planilha é convertida célula a
    célula (valores, mesclagens, fontes, bordas, larguras, moeda, orientação), então nunca diz outra coisa. Fórmula vira o valor calculado."""
    import io, zipfile
    from orcamento import db, pacote, pacote_pdf
    pid = _projeto(cliente)
    p = _plano(pid)
    wb = pacote.montar_planilha(p)
    pags = {x['titulo']: x for x in pacote_pdf.paginas_html(wb)}
    assert list(pags) == wb.sheetnames and [x['paisagem'] for x in pags.values()] == [False, True, True, False, True]
    assert pags['Comparativo de Preço']['escala'] == 0.74 and pags['Plano de Aplicação']['escala'] == 1.0 and pags['Cronograma fisico-financeiro']['escala'] == 1.0   # a escala da planilha; o que não couber na página é reduzido
    plano = pags['Plano de Aplicação']['html']
    for trecho in ('PLANO DE APLICAÇÃO', 'Concedente<br> (SJC)', 'Psicólogo (recibo ou MEI) - 100h mensal', 'Café 500g (2 unidades x R$10,00)', '<span>R$</span><span>52.280,00</span>',
                   'colspan="6"', 'rowspan="3"', "font-family:'Aptos Narrow'", 'border-left:2px solid #000', 'background:#000000', 'color:#FF0000'):
        assert trecho in plano, trecho
    assert '=SUM' not in plano and '=ROUND' not in plano                                                        # nenhuma fórmula aparece: só o valor
    crono = pags['Cronograma fisico-financeiro']['html']
    assert 'Psicólogo (recibo ou MEI) - 100h mensal' in crono and '<span>R$</span><span>3.020,00</span>' in crono and '14º MÊS' in crono and "'Plano de Aplicação'!" not in crono
    assert '-R$' not in crono                                                                                    # idem: o "teto − total" não é impresso
    comp = pags['Comparativo de Preço']['html']
    assert comp.count('<thead>') == 1 and comp.index('Valor Médio') < comp.index('</thead>') < comp.index('Psicólogo')   # o cabeçalho repete em cada página
    assert 'ALFA - CNPJ: 11.222.333/0001-81' in comp and 'class="g"' in comp
    assert pacote_pdf.moeda(1234.5) == ('R$', '1.234,50') and pacote_pdf.moeda(0) == ('R$', '-') and pacote_pdf.moeda(-42280) == ('-R$', '42.280,00')
    # baixar pela tela (aqui com o PDF de teste no lugar do navegador) e dentro do pacote
    r = cliente.get(f'/p/{pid}/planilhas-pdf')
    assert r.status_code == 200 and r.headers['content-type'] == 'application/pdf' and r.content[:5] == b'%PDF-'
    assert 'href="/p/%d/planilhas-pdf"' % pid in cliente.get(f'/p/{pid}').text
    z = zipfile.ZipFile(io.BytesIO(cliente.get(f'/p/{pid}/pacote').content))
    nomes = [n.split('/', 1)[1] for n in z.namelist()]
    assert nomes[:2] == ['Plano de Aplicação e Comparativo de Preço.xlsx', 'Plano de Aplicação e Comparativo de Preço.pdf']
    assert '"Plano de Aplicação e Comparativo de Preço.pdf": as mesmas planilhas' in z.read([n for n in z.namelist() if n.endswith('LEIA-ME.txt')][0]).decode('utf-8')


def test_pacote_sai_mesmo_sem_o_navegador(cliente, monkeypatch):
    """Sem o navegador do sistema, o pacote não deixa de sair: vai sem o PDF das planilhas, e o LEIA-ME avisa."""
    import io, zipfile
    from orcamento import pacote_pdf
    pid = _projeto(cliente); _plano(pid)

    def falha(wb):
        raise RuntimeError('sem navegador')
    monkeypatch.setattr(pacote_pdf, 'pdf_da_planilha', falha)
    z = zipfile.ZipFile(io.BytesIO(cliente.get(f'/p/{pid}/pacote').content))
    nomes = [n.split('/', 1)[1] for n in z.namelist()]
    assert nomes[0].endswith('.xlsx') and not any(n.endswith('Preço.pdf') for n in nomes)
    assert 'O PDF das planilhas não pôde ser gerado agora (RuntimeError)' in z.read([n for n in z.namelist() if n.endswith('LEIA-ME.txt')][0]).decode('utf-8')
    r = cliente.get(f'/p/{pid}/planilhas-pdf', follow_redirects=True)
    assert 'Não foi possível gerar o PDF agora (RuntimeError)' in r.text


@pytest.mark.navegador
def test_pdf_das_planilhas_impresso_de_verdade(cliente):
    """A impressão de verdade, pelo navegador do sistema: uma página por aba (o Comparativo pode ter mais), na orientação da planilha."""
    import pymupdf
    from orcamento import pacote
    pid = _projeto(cliente)
    p = _plano(pid)
    try:
        dados = pacote.planilhas_pdf(p)
    except Exception as e:   # computador sem o navegador do sistema instalado
        pytest.skip(f'navegador indisponível: {type(e).__name__}')
    d = pymupdf.open(stream=dados, filetype='pdf')
    textos = [' '.join(pg.get_text().split()) for pg in d]
    assert len(d) == 5 and [pg.rect.width > pg.rect.height for pg in d] == [False, True, True, False, True]
    assert textos[0].startswith('PLANO DE APLICAÇÃO') and 'R$ 52.280,00' in textos[0] and 'Notebook 15 polegadas (1 unidade x R$3.000,00)' in textos[0]
    assert textos[1].startswith('CRONOGRAMA FÍSICO-FINANCEIRO') and '14º MÊS' in textos[1] and 'TOTAL AGRUPADO' in textos[1]
    assert textos[2].startswith('ETAPAS E FASES') and '3º ao 12º mês' in textos[2] and textos[3].startswith('CRONOGRAMA DE DESEMBOLSO')
    assert textos[4].startswith('COMPARATIVO DE PREÇOS') and 'ALFA - CNPJ: 11.222.333/0001-81' in textos[4]


def test_ia_sugere_itens_e_quantidades_e_a_pessoa_decide(cliente, monkeypatch):
    """05/10/2026: na rubrica de produtos, a IA sugere os itens e a quantidade por mês a partir do que a OSC descreve. Nada é gravado antes de a
    pessoa aceitar; a IA não sugere preço (preço só vem de pesquisa), a marca sugerida é tirada e o que já está na rubrica não se repete."""
    import json
    from orcamento import db, ia, servico
    from orcamento.modelo import RubricaMaterial, Subitem
    pid = _projeto(cliente)
    p = db.carregar(pid)[0]
    p.rubricas = [RubricaMaterial(item=1, descricao='Alimentação', meses=10, regra='mercado', subitens=[Subitem(descricao='Café', especificacao='500g', qtd=2)])]
    db.salvar(pid, p)
    v0 = db.carregar(pid)[1]
    pag = cliente.get(f'/p/{pid}/mat/1').text
    assert 'Peça à IA uma lista de itens e quantidades' in pag and 'name="para_que"' in pag and 'O que a IA sugere' not in pag
    pedidos = []

    def falsa(rubrica, para_que, ja_tem, projeto_id=None):
        pedidos.append((rubrica, para_que, ja_tem))
        return [dict(descricao='café', especificacao='500g', quantidade=3, motivo='já existe'),
                dict(descricao='Leite Integral Italac', especificacao='1l', quantidade=20, motivo='30 pessoas x 8 encontros: 1 caixa para 12 copos', preco=5.39),
                dict(descricao='Suco de Uva', especificacao='1 litro', quantidade='8', motivo='um litro para 4 pessoas'),
                dict(descricao='Geladeira', especificacao='', quantidade=0), dict(descricao='Pão de Forma', especificacao='500g', quantidade=1200), 'texto solto']
    monkeypatch.setattr(ia, 'disponivel', lambda: True)
    monkeypatch.setattr(ia, 'sugerir_itens', falsa)
    assert 'descreva a atividade com um pouco mais de detalhe' in cliente.post(f'/p/{pid}/mat/1/sugerir', data={'para_que': 'lanche'}).text and not pedidos
    pag = cliente.post(f'/p/{pid}/mat/1/sugerir', data={'para_que': 'Lanche para 30 adolescentes, em 2 encontros por semana.'}).text
    assert pedidos == [('Alimentação', 'Lanche para 30 adolescentes, em 2 encontros por semana.', ['Café 500g'])]            # só isso vai para a IA
    assert 'O que a IA sugere para esta rubrica' in pag and 'Nada foi gravado' in pag and 'os preços vêm da pesquisa nas lojas' in pag
    assert 'Leite Integral' in pag and 'Italac' not in pag and 'Suco de Uva' in pag and '30 pessoas x 8 encontros' in pag
    for fora in ('Geladeira', 'já existe', '5,39', '5.39', '1200'):                                                           # quantidade inválida, repetido, preço: ficam de fora
        assert fora not in pag.split('id="sugestoes"')[1].split('</section>')[0], fora
    assert db.carregar(pid)[1] == v0                                                                                          # nada gravado: é só uma sugestão
    dados = json.loads(re.search(r'name="dados" value="([^"]*)"', pag).group(1).replace('&#34;', '"').replace('&amp;', '&'))
    assert [(x['descricao'], x['especificacao'], x['quantidade']) for x in dados] == [('Leite Integral', '1L', 20), ('Suco de Uva', '1L', 8)]
    # a pessoa aceita só o leite, com outra quantidade; o que volta da tela é conferido de novo (não se aceita item adulterado)
    adulterado = json.dumps(dados + [dict(descricao='Notebook Dell', especificacao='', quantidade=5000)], ensure_ascii=False)
    r = cliente.post(f'/p/{pid}/mat/1/sugerir/aplicar', data={'dados': adulterado, 'sugestao': ['0', '2'], 'qtd0': '12', 'qtd2': '5000'}, follow_redirects=False)
    assert r.status_code == 303 and '#t-itens' in r.headers['location']
    p = db.carregar(pid)[0]
    assert [(s.descricao, s.especificacao, s.qtd, s.valor_plano, s.precos) for s in p.rubricas[0].subitens] == \
        [('Café', '500g', 2, None, [None, None, None]), ('Leite Integral', '1L', 12, None, [None, None, None])]                   # sem preço: vem da pesquisa
    assert 'sugerido(s) pela IA e aceito(s) pela OSC' in db.historico(pid)[0][0]['motivo']
    r = cliente.post(f'/p/{pid}/mat/1/sugerir/aplicar', data={'dados': json.dumps(dados, ensure_ascii=False)}, follow_redirects=True)
    assert 'nenhuma sugestão estava marcada' in r.text and len(db.carregar(pid)[0].rubricas[0].subitens) == 2
    # sem a IA, a tela diz o que fazer; e resposta vazia ou falha não quebram
    monkeypatch.setattr(ia, 'sugerir_itens', lambda *a, **k: None)
    assert 'Não foi possível consultar a IA agora' in cliente.post(f'/p/{pid}/mat/1/sugerir', data={'para_que': 'Lanche para 30 adolescentes, 2 vezes por semana'}).text
    monkeypatch.setattr(ia, 'sugerir_itens', lambda *a, **k: [])
    assert 'A IA não sugeriu nenhum item novo' in cliente.post(f'/p/{pid}/mat/1/sugerir', data={'para_que': 'Lanche para 30 adolescentes, 2 vezes por semana'}).text
    monkeypatch.setattr(ia, 'disponivel', lambda: False)
    assert 'a IA gratuita não está configurada' in cliente.post(f'/p/{pid}/mat/1/sugerir', data={'para_que': 'Lanche para 30 adolescentes, 2 vezes por semana'}).text
    assert servico.sugestoes_conferidas(p.rubricas[0], None) == [] and servico.sugestoes_conferidas(p.rubricas[0], [dict(descricao='', quantidade=1)]) == []
