"""v0.3: título exato das vagas, tarefas com progresso/cancelamento, base da Receita (homônimos) e grade com fornecedores por item."""
import os
import sqlite3
import time

import pytest

from orcamento.vagas import titulo_exato, avaliar


@pytest.mark.parametrize('titulo, cargo, esperado', [
    ('Coordenador(a) de Projetos', 'Coordenador de Projetos', True), ('Coordenadora de Projeto', 'Coordenador de Projetos', True),
    ('COORDENADOR DE PROJETOS', 'Coordenador de Projetos', True), ('Coordenador/a de Projetos', 'Coordenador de Projetos', True),
    ('Coordenador de Projetos de TI', 'Coordenador de Projetos', False), ('Coordenador de Projetos Sociais', 'Coordenador de Projetos', False),
    ('Coordenador de Projetos - São Paulo', 'Coordenador de Projetos', True), ('Auxiliar de Serviço Geral', 'Auxiliar de serviços gerais', True),
    ('Psicóloga', 'Psicólogo', True), ('Psicólogo Clínico', 'Psicólogo', False), ('Orientadora Socioeducativa', 'Orientador socioeducativo', True),
    ('Designer Gráfico Pleno', 'Designer Gráfico', False), ('Auxiliar Administrativa', 'Auxiliar Administrativo', True),
])
def test_titulo_exato(titulo, cargo, esperado):
    assert titulo_exato(titulo, cargo) == esperado


def test_agregador_nao_vale():
    v = dict(titulo='Assistente Social', empresa='OEmprego - Vagas Brasil', faixa_min=300000, faixa_max=300000, unidade='MONTH')
    assert 'agregador' in avaliar(v, 'Assistente Social')


@pytest.fixture
def dados(tmp_path, monkeypatch):
    import orcamento.db as dbm
    monkeypatch.setattr(dbm, 'PASTA_DADOS', str(tmp_path))
    return tmp_path


def test_tarefa_progresso_e_cancelamento(dados):
    from orcamento import tarefas

    def trabalho(ctx):
        ctx.fonte('Loja A', 'ok'); ctx.fonte('Loja B', 'falhou', 'sem resposta em 75 s')
        ctx.aviso('RESULTADO INCOMPLETO — Loja B')
        for i in range(50):
            ctx.progresso(i * 2, f'passo {i}'); time.sleep(0.05)
        return dict(ok=True, ir_para='/x')
    tid = tarefas.iniciar('teste', 'tarefa de teste', trabalho)
    time.sleep(0.6)
    t = tarefas.ler(tid)
    assert t['estado'] == 'rodando' and t['fontes']['Loja B']['estado'] == 'falhou' and t['avisos']
    tarefas.cancelar(tid)
    for _ in range(50):
        if tarefas.ler(tid)['estado'] != 'rodando':
            break
        time.sleep(0.1)
    assert tarefas.ler(tid)['estado'] == 'cancelada'
    tid2 = tarefas.iniciar('teste', 'termina', lambda ctx: dict(ir_para='/p/1'), projeto_id=1)   # com projeto: registra no histórico
    for _ in range(50):
        if tarefas.ler(tid2)['estado'] == 'concluída':
            break
        time.sleep(0.1)
    t2 = tarefas.ler(tid2)
    assert t2['estado'] == 'concluída' and t2['progresso'] == 100 and t2['resultado']['ir_para'] == '/p/1'


def _base_exemplo(caminho):
    """Mini base no formato do índice: 2 empresas 'Confiança RH' em Niterói e 1 empresa de nome único."""
    c = sqlite3.connect(caminho)
    c.executescript("""CREATE TABLE empresa(basico TEXT PRIMARY KEY, razao TEXT);
        CREATE TABLE estab(cnpj TEXT PRIMARY KEY, basico TEXT, matriz INTEGER, fantasia TEXT, uf TEXT, municipio TEXT, cnae TEXT);
        CREATE TABLE municipio(codigo TEXT PRIMARY KEY, nome TEXT); CREATE TABLE meta(chave TEXT PRIMARY KEY, valor TEXT);
        CREATE VIRTUAL TABLE nomes USING fts5(nome, cnpj UNINDEXED, tokenize='unicode61 remove_diacritics 2');
        INSERT INTO municipio VALUES ('5865','NITEROI'),('7107','SAO PAULO');
        INSERT INTO empresa VALUES ('03472547','CONFIANCA RH, CONSULTORIA E PRESTADORA DE SERVICOS LTDA'),('12802628','CONFIANCA SERVICOS DE ASSESSORIA LTDA'),
                                   ('13312641','DNA FACILITIES LTDA');
        INSERT INTO estab VALUES ('03472547000188','03472547',1,'','RJ','5865','7820500'),('12802628000190','12802628',1,'CONFIANCA RH','RJ','5865','7820500'),
                                 ('13312641000123','13312641',1,'','SP','7107','8111700');
        INSERT INTO meta VALUES ('pasta_origem','2026-09');
        INSERT INTO nomes(nome, cnpj) SELECT COALESCE(e.fantasia,'') || ' | ' || COALESCE(m.razao,''), e.cnpj FROM estab e LEFT JOIN empresa m ON m.basico=e.basico;""")
    c.commit(); c.close()


def test_base_da_receita_homonimos_e_nome_unico(dados, monkeypatch):
    from orcamento import cnpj_base, cnpj_busca
    arq = os.path.join(dados, 'cnpj_teste.sqlite'); _base_exemplo(arq)
    monkeypatch.setattr(cnpj_base, 'BASE', arq)
    assert cnpj_busca.cnpj_pela_base('DNA Facilities')['cnpj'] == '13312641000123'
    assert cnpj_busca.cnpj_pela_base('Confiança RH') is None                          # 2 empresas com o nome: não escolhe
    r = cnpj_busca.conferir_homonimos(dict(status='🟢', cnpj='12802628000190', motivo='nome confere', fonte='Yahoo'), 'Confiança RH')
    assert r['status'] == '🟡' and 'mais 1 empresa' in r['motivo']
    assert cnpj_busca.conferir_local(dict(status='🟢', cnpj='1', municipio='EXTERIOR', uf='EX', motivo='x'), 'Goiás', None)['status'] == '🔴'


def test_sede_em_outro_estado_so_e_duvida_quando_ha_homonima(dados, monkeypatch):
    """Pedido da OSC (06/10/2026): muitas vagas ficavam "em dúvida" por causa do nome da empresa. Medido no banco: 10 de 31 eram só "CNPJ em
    outro estado", sem haver outra empresa com o nome. A vaga diz onde é o TRABALHO; a sede pode ser em outro lugar."""
    import asyncio
    from orcamento import cnpj_base, cnpj_busca
    arq = os.path.join(dados, 'cnpj_teste.sqlite'); _base_exemplo(arq)
    monkeypatch.setattr(cnpj_base, 'BASE', arq)
    # empresa única com esse nome no Brasil, sede em SP, vaga no Ceará: é ela
    unica = dict(status='🟢', cnpj='13312641000123', municipio='SAO PAULO', uf='SP', motivo='nome confere', fonte='Yahoo')
    r = cnpj_busca.conferir_local(unica, 'Fortaleza', 'CE', 'DNA Facilities')
    assert r['status'] == '🟢' and 'sede em SP, vaga em CE' in r['motivo'] and 'única empresa ativa com esse nome' in r['motivo']
    # há outra empresa com o mesmo nome: continua em dúvida
    r = cnpj_busca.conferir_local(dict(status='🟢', cnpj='12802628000190', municipio='NITEROI', uf='RJ', motivo='nome confere', fonte='Yahoo'), 'São Paulo', 'SP', 'Confiança RH')
    assert r['status'] == '🟡' and cnpj_busca.OUTRA_UF in r['motivo']
    # CNPJ publicado no site oficial da empresa (ou escrito na vaga): a sede em outro estado não é dúvida, mesmo havendo homônimas
    r = cnpj_busca.conferir_local(dict(status='🟢', cnpj='12802628000190', municipio='NITEROI', uf='RJ', motivo='CNPJ publicado no site oficial', fonte='site oficial https://x'),
                                  'São Paulo', 'SP', 'Confiança RH')
    assert r['status'] == '🟢' and 'sede em RJ, vaga em SP' in r['motivo']
    assert cnpj_busca.conferir_local(unica, 'São Paulo', 'SP', 'DNA Facilities') == unica                 # mesmo estado: nada muda
    assert cnpj_busca.conferir_local(unica, 'Fortaleza', 'CE')['status'] == '🟡'                            # sem o nome não dá para conferir as homônimas
    # dúvida guardada (memória de 30 dias) só por causa do estado: é conferida de novo com a regra atual, sem consultar a internet
    from orcamento import db
    chave = 'dna facilities|fortaleza|ce'
    with db.conectar() as c:
        c.execute('INSERT OR REPLACE INTO empresa_cnpj (chave, consultado_em, json) VALUES (?,?,?)',
                  (chave, db.agora(), '{"status": "🟡", "cnpj": "13312641000123", "uf": "SP", "municipio": "SAO PAULO", "fonte": "Yahoo", '
                                      '"motivo": "nome confere; CNPJ em SP, vaga em CE: pode ser filial ou homônima"}'))
    r = asyncio.run(cnpj_busca.cnpj_do_empregador(None, None, 'DNA Facilities', 'Fortaleza', 'CE'))
    assert r['status'] == '🟢' and r['cnpj'] == '13312641000123' and r['motivo'].startswith('nome confere; sede em SP, vaga em CE')
    # CNPJ escrito no texto da própria vaga: vale (ativo na base e o nome confere), antes de qualquer busca
    from orcamento.vagas import cnpjs_no_texto
    assert cnpjs_no_texto('<p>Empresa DNA Facilities, CNPJ 13.312.641/0001-23. Contato...</p> CNPJ: 13312641000123') == ['13312641000123']
    r = asyncio.run(cnpj_busca.cnpj_do_empregador(None, None, 'DNA Facilities', 'Rio de Janeiro', 'RJ', cnpjs_texto=['13312641000123']))
    assert r['status'] == '🟢' and r['fonte'] == 'texto da vaga' and 'escrito no texto da própria vaga' in r['motivo'] and 'sede em SP, vaga em RJ' in r['motivo']
    assert asyncio.run(cnpj_busca.cnpj_do_texto(None, 'Outra Empresa Qualquer', ['13312641000123'])) is None   # o CNPJ escrito é de outra empresa: não vale
    # para a OSC escolher com um clique: as empresas com esse nome, as do estado da vaga primeiro
    cs = cnpj_busca.candidatos_da_base('Confiança RH', 'Niterói', 'RJ')
    assert [(c['cnpj'], c['municipio'], c['na_cidade']) for c in cs] == [('12802628000190', 'NITEROI', True), ('03472547000188', 'NITEROI', True)]   # o nome idêntico antes do nome contido
    assert cnpj_busca.candidatos_da_base('Empresa Que Não Existe', 'Niterói', 'RJ') == []


def test_nova_fonte_de_vagas_titulo_no_endereco_e_milhar_escrito_com_ponto():
    """Pedido da OSC (08/10/2026): melhorar a busca de vagas. Sondados 8 sites; entrou o Empregos.com.br (título, empresa e salário legíveis, sem
    verificação humana). Nele o link da vaga só diz "Mais detalhes" (o título está no endereço) e o salário de R$ 2.200 vem como "2.2"."""
    from orcamento import vagas as V
    nomes = [f[0] for f in V.fontes('Auxiliar Administrativo')]
    assert 'Empregos.com.br' in nomes and 'https://www.empregos.com.br/vagas/auxiliar-administrativo' in [f[1] for f in V.fontes('Auxiliar Administrativo')]
    assert 'Empregos.com.br' not in [f[0] for f in V.fontes('Auxiliar Administrativo', profundo=True)]
    assert V.titulo_do_endereco('https://www.empregos.com.br/vaga/11891982/auxiliar-administrativo-em-atibaia-sp') == 'auxiliar administrativo em atibaia sp'
    assert V.titulo_confere(V.titulo_do_endereco('https://www.empregos.com.br/vaga/11891982/auxiliar-administrativo-em-atibaia-sp'), 'Auxiliar Administrativo')
    assert not V.titulo_confere(V.titulo_do_endereco('https://www.empregos.com.br/vaga/11891982/motorista-em-atibaia-sp'), 'Auxiliar Administrativo')
    # o valor escrito do jeito americano também confirma o salário na página guardada
    assert V.salario_no_texto('Remuneração R$2,295.00 Publicado há 1 dia', 229500) and V.salario_no_texto('Remuneração R$ 2.295,00', 229500)
    assert not V.salario_no_texto('R$2,295.50', 229500) and not V.salario_no_texto('R$ 12,295.00', 229500)
    # "2.295" nos dados da vaga: é R$ 2.295,00 SÓ se a página mostrar esse valor; senão, a vaga fica sem salário
    v = dict(titulo='Auxiliar Administrativo', empresa='Empresa Exemplo', faixa_min=230, faixa_max=230, unidade='MONTH', faixa_bruta=[2.295, 2.295])
    V.milhar_com_ponto(v, 'Nº de vagas 1 Remuneração R$2,295.00 Publicado há 1 dia')
    assert (v['faixa_min'], v['faixa_max'], v['salario_da']) == (229500, 229500, 'página (milhar)') and avaliar(v, 'Auxiliar administrativo') is None
    v = dict(titulo='Auxiliar Administrativo', empresa='Empresa Exemplo', faixa_min=220, faixa_max=220, unidade='MONTH', faixa_bruta=[2.2, 2.2])
    V.milhar_com_ponto(v, 'Página sem o valor por extenso')
    assert v['faixa_min'] is None and avaliar(v, 'Auxiliar administrativo') == 'sem salário informado'
    # salário mensal de menos de R$ 100 nunca passa (com o fim do mínimo de R$ 1.000, um erro de leitura viraria a vaga "mais barata")
    assert 'menos de R$ 100' in avaliar(dict(titulo='Auxiliar Administrativo', empresa='Empresa Exemplo', faixa_min=220, faixa_max=220, unidade='MONTH'), 'Auxiliar administrativo')
    v = dict(titulo='Auxiliar Administrativo', empresa='Empresa Exemplo', faixa_min=180000, faixa_max=180000, unidade='MONTH', faixa_bruta=[1800, 1800])
    assert V.milhar_com_ponto(v, 'qualquer texto')['faixa_min'] == 180000                                 # salário normal: nada muda


def test_grade_com_fornecedor_por_item(dados, tmp_path):
    from openpyxl import load_workbook
    from orcamento.modelo import Projeto, RubricaMaterial, Subitem, Fonte
    from orcamento.exportar import exportar
    fs = [Fonte(nome=n, cnpj=c, data_pesquisa='2026-09-27') for n, c in
          [('ATACADAO S.A.', '75.315.333/0001-09'), ('CARREFOUR COMERCIO E INDUSTRIA LTDA', '45.543.915/0001-81'), ('TENDA ATACADO SA', '01.157.555/0011-86')]]
    s = Subitem(descricao='Maionese 250g', qtd=10, precos=[858, 999, 815], fontes=fs, valor_plano=890)
    p = Projeto(nome='t', teto=8900, rubricas=[RubricaMaterial(item=1, descricao='Alimentação', meses=1, subitens=[s])])
    arq = os.path.join(tmp_path, 'g.xlsx'); exportar(p, [], {}, arq)
    ws = load_workbook(arq)['Grade Comparativa']
    textos = [str(c.value) for row in ws.iter_rows() for c in row if c.value]
    assert any('ATACADAO S.A. - CNPJ: 75.315.333/0001-09' in t for t in textos) and any('TENDA ATACADO SA' in t for t in textos)
    assert any('fornecedores por subitem' in t for t in textos)


def test_banco_de_vagas_cria_a_pasta(dados):
    from orcamento import vagas as V
    v = dict(titulo='Auxiliar de Serviços Gerais', empresa='Adecco', url='https://x/vaga/1', plataforma='InfoJobs', faixa_min=180000,
             faixa_max=180000, unidade='MONTH', cidade='São Paulo', uf='SP')
    V.guardar_no_banco('Auxiliar de serviços gerais', v, dict(status='🟢', cnpj='12345678000199', razao='ADECCO', motivo='teste'), b'%PDF-1.4 teste')
    assert os.path.exists(os.path.join(dados, 'banco_vagas'))
    assert any(f.endswith('_adecco.pdf') for f in os.listdir(os.path.join(dados, 'banco_vagas')))


@pytest.mark.parametrize('texto, esperado', [
    ('Função: SERVENTE/COPEIRA Salário: R$ 1.727.27 + 131,00 Gratificação', [172727]),
    ('<p>Salário base: R$ 2.000,00</p>', [200000]),
    ('Salário de R$ 1800 + VT', [180000]),
    ('Benefícios: vale alimentação R$ 900,00', []),
])
def test_salario_no_texto_da_vaga(texto, esperado):
    from orcamento.vagas import salarios_no_texto
    assert salarios_no_texto(texto) == esperado


def test_salario_do_texto_da_vaga_e_o_que_vale():
    """Decisão da OSC (06/10/2026): o texto da vaga pode citar o salário e, se citar, é o do texto que vale (antes a vaga era recusada)."""
    from orcamento.vagas import salarios_no_texto
    v = dict(titulo='Auxiliar de Serviços Gerais', empresa='Orsegups', faixa_min=200000, faixa_max=200000, unidade='MONTH', salarios_texto=[172727])
    assert avaliar(v, 'Auxiliar de serviços gerais') is None
    assert (v['faixa_min'], v['faixa_max'], v['salario_da']) == (172727, 172727, 'texto')
    v2 = dict(v, faixa_min=200000, faixa_max=200000, salarios_texto=[200000]); v2.pop('salario_da')
    assert avaliar(v2, 'Auxiliar de serviços gerais') is None and 'salario_da' not in v2            # o texto diz o mesmo: nada muda
    v3 = dict(v, faixa_min=None, faixa_max=None, salarios_texto=[250000, 180000])                  # o site não informa; o texto cita dois valores
    assert avaliar(v3, 'Auxiliar de serviços gerais') is None and (v3['faixa_min'], v3['faixa_max']) == (180000, 250000)
    # qualquer salário mensal vale: não há mais o mínimo de R$ 1.000
    assert avaliar(dict(titulo='Auxiliar de Serviços Gerais', empresa='Orsegups', faixa_min=80000, faixa_max=80000, unidade='MONTH'), 'Auxiliar de serviços gerais') is None
    # valor por hora, por dia ou por aula escrito no texto não é salário mensal
    assert salarios_no_texto('Salário: R$ 1.727,27 + benefícios. Salário de R$ 25,00 por hora. Salário R$ 120,00/dia. salário R$ 800') == [172727, 80000]
    v4 = dict(v, faixa_min=200000, faixa_max=200000, salarios_texto=[172727], a_combinar=True)
    assert 'a combinar' in avaliar(v4, 'Auxiliar de serviços gerais') and v4['faixa_min'] == 200000   # "a combinar" continua fora
