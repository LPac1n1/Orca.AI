"""Exporta a Grade Comparativa, o Plano de Aplicação, a memória de cálculo, a verificação e a lista de CNPJs
num Excel com FÓRMULAS (médias e totais conferíveis célula a célula), no layout usado nos processos da SEJC."""
import datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .modelo import Projeto, RubricaRH, RubricaMaterial, fontes_do_subitem, descricao_completa
from .calculo import media_rh, media_subitem, mensal_maximo_rh, valores_pesquisa, total_projeto
from .regras import REGRAS, divisor_horas, valor_hora, brl, cnpj_formatar, enquadrar
from .cnpj import URL_COMPROVANTE

MOEDA = '"R$" #,##0.00'
FINO = Side(style='thin', color='999999')
BORDA = Border(left=FINO, right=FINO, top=FINO, bottom=FINO)
CAB = PatternFill('solid', fgColor='D9E1F2')
RUB = PatternFill('solid', fgColor='F2F2F2')
ALERTA = {'erro': PatternFill('solid', fgColor='F8CBAD'), 'atencao': PatternFill('solid', fgColor='FFE699'), 'revisado': PatternFill('solid', fgColor='DDEBF7'),
          'info': PatternFill('solid', fgColor='E2EFDA')}
GRAVIDADES = ('erro', 'atencao', 'revisado', 'info')


def _desc_rh(r: RubricaRH):
    reg = {'RECIBO_MEI': 'recibo ou MEI', 'PJ': 'PJ', 'CLT': 'CLT'}[r.regime]
    return f'{r.cargo} ({reg}) - {r.horas_mes}h mensal'


def _qtd_rh(r: RubricaRH):
    return max(1, r.quantidade or 1)


def _cel(ws, ref, v, fmt=None, bold=False, fill=None, wrap=False):
    c = ws[ref]; c.value = v; c.border = BORDA
    if fmt: c.number_format = fmt
    if bold: c.font = Font(bold=True)
    if fill: c.fill = fill
    c.alignment = Alignment(wrap_text=wrap, vertical='center')
    return c


def _fornecedor(f):
    return f'{f.nome} - CNPJ: {cnpj_formatar(f.cnpj)}' if f.cnpj else f.nome


def grade(wb, p: Projeto):
    ws = wb.active; ws.title = 'Grade Comparativa'
    ws.merge_cells('A1:M1'); _cel(ws, 'A1', 'COMPARATIVO DE PREÇOS', bold=True, fill=CAB).alignment = Alignment(horizontal='center')
    for col, txt in (('D', 'ORÇAMENTO 1'), ('G', 'ORÇAMENTO 2'), ('J', 'ORÇAMENTO 3')):
        ws.merge_cells(f'{col}2:{get_column_letter(ws[col + "2"].column + 2)}2'); _cel(ws, f'{col}2', txt, bold=True, fill=CAB)
    for i, h in enumerate(['ITEM', 'DESCRIÇÃO DO ITEM', 'QTDE', 'PREÇO UNITÁRIO', 'PREÇO TOTAL', 'FORNECEDOR / CNPJ',
                           'PREÇO UNITÁRIO', 'PREÇO TOTAL', 'FORNECEDOR / CNPJ', 'PREÇO UNITÁRIO', 'PREÇO TOTAL', 'FORNECEDOR / CNPJ', 'VALOR MÉDIO'], 1):
        _cel(ws, f'{get_column_letter(i)}3', h, bold=True, fill=CAB, wrap=True)
    lin = 4
    for r in p.rubricas:
        if isinstance(r, RubricaRH):
            _cel(ws, f'A{lin}', r.item); _cel(ws, f'B{lin}', _desc_rh(r), wrap=True); _cel(ws, f'C{lin}', _qtd_rh(r))
            for k, (col, q) in enumerate(zip(('D', 'G', 'J'), r.pesquisas)):
                tot = get_column_letter(ws[col + '1'].column + 1); forn = get_column_letter(ws[col + '1'].column + 2)
                _cel(ws, f'{col}{lin}', (q.valor if q.valor is not None else q.faixa_min) / 100, MOEDA)
                _cel(ws, f'{tot}{lin}', f'={col}{lin}*C{lin}', MOEDA); _cel(ws, f'{forn}{lin}', _fornecedor(q), wrap=True)
            _cel(ws, f'M{lin}', f'=ROUND(AVERAGE(D{lin},G{lin},J{lin}),2)', MOEDA, bold=True)
            lin += 1
        else:
            cab = lin; lin += 1; primeira = lin
            for s in r.subitens:
                _cel(ws, f'B{lin}', descricao_completa(s), wrap=True); _cel(ws, f'C{lin}', s.qtd)
                fs = fontes_do_subitem(r, s)
                for k, col in enumerate(('D', 'G', 'J')):
                    tot = get_column_letter(ws[col + '1'].column + 1); forn = get_column_letter(ws[col + '1'].column + 2)
                    _cel(ws, f'{col}{lin}', (s.precos[k] or 0) / 100, MOEDA); _cel(ws, f'{tot}{lin}', f'={col}{lin}*C{lin}', MOEDA)
                    _cel(ws, f'{forn}{lin}', _fornecedor(fs[k]) if k < len(fs) else '', wrap=True)
                _cel(ws, f'M{lin}', f'=ROUND(AVERAGE(D{lin},G{lin},J{lin}),2)', MOEDA)
                lin += 1
            ult = lin - 1
            _cel(ws, f'A{cab}', r.item, fill=RUB, bold=True); _cel(ws, f'B{cab}', r.descricao, fill=RUB, bold=True, wrap=True); _cel(ws, f'C{cab}', '', fill=RUB)
            for k, col in enumerate(('D', 'G', 'J')):
                tot = get_column_letter(ws[col + '1'].column + 1); forn = get_column_letter(ws[col + '1'].column + 2)
                _cel(ws, f'{col}{cab}', f'=SUM({col}{primeira}:{col}{ult})', MOEDA, fill=RUB)
                _cel(ws, f'{tot}{cab}', f'=SUM({tot}{primeira}:{tot}{ult})', MOEDA, bold=True, fill=RUB)
                por_item = all(len(s.fontes) == 3 for s in r.subitens) and r.subitens
                _cel(ws, f'{forn}{cab}', 'fornecedores por subitem (abaixo)' if por_item else (_fornecedor(r.fontes[k]) if k < len(r.fontes) else ''), fill=RUB, wrap=True)
            _cel(ws, f'M{cab}', f'=ROUND(AVERAGE(E{cab},H{cab},K{cab}),2)', MOEDA, bold=True, fill=RUB)
    for col, w in zip('ABCDEFGHIJKLM', (6, 38, 6, 13, 13, 30, 13, 13, 30, 13, 13, 30, 14)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = 'C4'


def plano(wb, p: Projeto):
    ws = wb.create_sheet('Plano de Aplicação')
    ws.merge_cells('A1:F1'); _cel(ws, 'A1', '15. PLANO DE APLICAÇÃO DOS RECURSOS FINANCEIROS', bold=True, fill=CAB)
    for i, h in enumerate(['Item', 'Descrição', 'Valor Unitário', 'Valor Total', 'Concedente (SJC)', 'Proponente (entidade)'], 1):
        _cel(ws, f'{get_column_letter(i)}2', h, bold=True, fill=CAB)
    lin = 3; totais = []
    for r in p.rubricas:
        if isinstance(r, RubricaRH):
            q, v = _qtd_rh(r), (r.valor_mensal_plano or 0)
            _cel(ws, f'A{lin}', r.item)
            if q > 1:   # vários profissionais no mesmo item: o valor unitário é o mensal do item (como nos subitens: quantidade x valor)
                _cel(ws, f'B{lin}', f'{_desc_rh(r)} ({q} profissionais x {brl(v).replace(" ", "")})', wrap=True)
                _cel(ws, f'C{lin}', f'={q}*{v / 100}', MOEDA)
            else:
                _cel(ws, f'B{lin}', _desc_rh(r), wrap=True); _cel(ws, f'C{lin}', v / 100, MOEDA)
            _cel(ws, f'D{lin}', f'=ROUND(C{lin}*{r.meses},2)', MOEDA)
            _cel(ws, f'E{lin}', f'=D{lin}', MOEDA); _cel(ws, f'F{lin}', '')
            totais.append(f'D{lin}'); lin += 1
        else:
            cab = lin; lin += 1; primeira = lin
            for s in r.subitens:
                v = (s.valor_plano or 0)
                _cel(ws, f'B{lin}', f'{descricao_completa(s)} ({s.qtd} unidade{"s" if s.qtd > 1 else ""} x {brl(v).replace(" ", "")})', wrap=True)
                _cel(ws, f'C{lin}', f'={s.qtd}*{v / 100}', MOEDA); lin += 1
            _cel(ws, f'A{cab}', r.item, bold=True, fill=RUB); _cel(ws, f'B{cab}', r.descricao, bold=True, fill=RUB)
            _cel(ws, f'C{cab}', f'=SUM(C{primeira}:C{lin - 1})', MOEDA, bold=True, fill=RUB)
            _cel(ws, f'D{cab}', f'=ROUND(C{cab}*{r.meses},2)', MOEDA, bold=True, fill=RUB)
            _cel(ws, f'E{cab}', f'=D{cab}', MOEDA, fill=RUB); _cel(ws, f'F{cab}', '', fill=RUB)
            totais.append(f'D{cab}')
    _cel(ws, f'B{lin}', 'Total', bold=True); _cel(ws, f'D{lin}', '=' + '+'.join(totais), MOEDA, bold=True)
    _cel(ws, f'E{lin}', f'=D{lin}', MOEDA, bold=True); _cel(ws, f'F{lin}', 0, MOEDA)
    _cel(ws, f'B{lin + 1}', 'Teto (recurso da emenda)'); _cel(ws, f'D{lin + 1}', p.teto / 100, MOEDA)
    _cel(ws, f'B{lin + 2}', 'Conferência'); _cel(ws, f'D{lin + 2}', f'=IF(ROUND(D{lin},2)=ROUND(D{lin + 1},2),"FECHA NO TETO","NÃO FECHA")', bold=True)
    for col, w in zip('ABCDEF', (6, 70, 15, 15, 16, 16)):
        ws.column_dimensions[col].width = w


def memoria(wb, p: Projeto):
    ws = wb.create_sheet('Memória de Cálculo')
    cab = ['Item', 'Cargo / subitem', 'Valores das 3 pesquisas', 'Média', 'Divisor de horas', 'Base', 'Valor-hora', 'Horas/mês',
           'Valor mensal/unitário', 'Meses', 'Total', 'Origem do valor']
    for i, h in enumerate(cab, 1):
        _cel(ws, f'{get_column_letter(i)}1', h, bold=True, fill=CAB, wrap=True)
    lin = 2
    for r in p.rubricas:
        if isinstance(r, RubricaRH):
            maxm, vh, div = mensal_maximo_rh(r, p.config)
            enq = enquadrar(r.cargo)
            base = f'{"jornada legal: " + str(enq["semanais"]) + "h/sem ÷ 6 × 30 (" + enq["base_legal"] + ")" if p.config.divisor_horas == "legal" else "divisor já aceito pela SEJC"}'
            vals = ' | '.join(brl(v) for v in valores_pesquisa(r))
            q = _qtd_rh(r)
            for i, v in enumerate([r.item, r.cargo + (f' ({q} profissionais)' if q > 1 else ''), vals, (media_rh(r) or 0) / 100, div, base, (vh or 0) / 100,
                                   r.horas_mes, (r.valor_mensal_plano or 0) / 100, r.meses, (r.valor_mensal_plano or 0) * r.meses * q / 100,
                                   'arredondar(média ÷ divisor; 2) × horas' + (f' × {q} profissionais' if q > 1 else '')], 1):
                _cel(ws, f'{get_column_letter(i)}{lin}', v, MOEDA if i in (4, 7, 9, 11) else None, wrap=i in (2, 3, 6))
            lin += 1
        else:
            for s in r.subitens:
                m = media_subitem(s)
                orig = 'média' if s.valor_plano == m else ('preço orçado (fonte %d)' % (s.precos.index(s.valor_plano) + 1) if s.valor_plano in s.precos else 'outro')
                for i, v in enumerate([r.item, descricao_completa(s), ' | '.join(brl(x) for x in s.precos), (m or 0) / 100, '', '', '', '',
                                       (s.valor_plano or 0) / 100, r.meses, (s.valor_plano or 0) * s.qtd * r.meses / 100, orig], 1):
                    _cel(ws, f'{get_column_letter(i)}{lin}', v, MOEDA if i in (4, 9, 11) else None, wrap=i in (2, 3))
                lin += 1
    for col, w in zip('ABCDEFGHIJKL', (6, 34, 40, 12, 10, 40, 11, 9, 14, 7, 14, 30)):
        ws.column_dimensions[col].width = w


def verificacao(wb, alertas):
    ws = wb.create_sheet('Verificação')
    for i, h in enumerate(['Gravidade', 'Regra', 'Descrição da regra', 'Origem', 'Item', 'Mensagem'], 1):
        _cel(ws, f'{get_column_letter(i)}1', h, bold=True, fill=CAB)
    for n, a in enumerate(sorted(alertas, key=lambda a: GRAVIDADES.index(a.gravidade)), 2):
        d = a.dict(); d['gravidade'] = {'erro': 'pendência', 'atencao': 'revisar', 'revisado': 'revisado pela OSC', 'info': 'informação'}[a.gravidade]
        for i, k in enumerate(['gravidade', 'regra', 'descricao_regra', 'origem', 'item', 'mensagem'], 1):
            _cel(ws, f'{get_column_letter(i)}{n}', d[k], fill=ALERTA[a.gravidade] if i == 1 else None, wrap=i in (3, 5, 6))
    for col, w in zip('ABCDEF', (16, 7, 45, 16, 40, 70)):
        ws.column_dimensions[col].width = w


def cnpjs(wb, p: Projeto, consultas: dict):
    ws = wb.create_sheet('CNPJs')
    for i, h in enumerate(['CNPJ', 'Empresa/Loja', 'Situação (API)', 'Consultado em', 'Fonte', 'Emitir comprovante oficial (Receita)',
                           'Comprovante oficial (arquivo)', 'Situação no comprovante', 'Emitido em'], 1):
        _cel(ws, f'{get_column_letter(i)}1', h, bold=True, fill=CAB)
    vistos = {}
    for r in p.rubricas:
        fs = list(r.pesquisas) if isinstance(r, RubricaRH) else list(r.fontes) + [f for s in r.subitens for f in fontes_do_subitem(r, s)]
        for f in fs:
            if f.cnpj or f.nome:
                vistos.setdefault(cnpj_formatar(f.cnpj), f.nome)
    for n, (c, nome) in enumerate(vistos.items(), 2):
        q = consultas.get(c, {})
        comp = (getattr(p, 'comprovantes_cnpj', None) or {}).get(c) or {}
        url = f'https://solucoes.receita.fazenda.gov.br/Servicos/cnpjreva/Cnpjreva_Solicitacao.asp?cnpj={"".join(ch for ch in c if ch.isdigit())}'
        for i, v in enumerate([c, nome, q.get('situacao') or 'não consultado', q.get('consultado_em', ''), q.get('fonte', ''), url,
                               (comp.get('arquivo') or 'FALTA — emitir e importar pelo sistema').split('/')[-1], comp.get('situacao') or '',
                               (comp.get('emitido_em') or '')[:10]], 1):
            cel = _cel(ws, f'{get_column_letter(i)}{n}', v)
            if i == 6:
                cel.hyperlink = url
    for col, w in zip('ABCDEFGHI', (20, 50, 16, 26, 14, 60, 44, 14, 12)):
        ws.column_dimensions[col].width = w


def exportar(p: Projeto, alertas, consultas_cnpj, caminho):
    wb = Workbook()
    grade(wb, p); plano(wb, p); memoria(wb, p); verificacao(wb, alertas); cnpjs(wb, p, consultas_cnpj)
    info = wb.create_sheet('Sobre')
    for i, (k, v) in enumerate([('Projeto', p.nome), ('Processo', p.processo), ('Proponente', p.proponente), ('Teto', brl(p.teto)),
                                ('Total do plano', brl(total_projeto(p))), ('Divisor de horas', p.config.divisor_horas),
                                ('Valores defensáveis', 'sim' if p.config.valores_defensaveis else 'não'),
                                ('Gerado em', dt.datetime.now().strftime('%d/%m/%Y %H:%M')),
                                ('Observação', 'Médias e totais são fórmulas: podem ser conferidos célula a célula.')], 1):
        info[f'A{i}'] = k; info[f'B{i}'] = v; info[f'A{i}'].font = Font(bold=True)
    info.column_dimensions['A'].width = 20; info.column_dimensions['B'].width = 80
    wb.save(caminho)
    return caminho
