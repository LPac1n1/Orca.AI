"""As planilhas do pacote em PDF (05/10/2026).

O PDF não é montado à parte: cada aba da MESMA planilha que o sistema grava (.xlsx) é convertida célula a célula — valores, mesclagens, fontes,
bordas, fundos, larguras, formato de moeda, orientação e escala da página —, e o navegador do sistema (o mesmo que guarda os comprovantes)
imprime cada aba. Assim o PDF e a planilha nunca dizem coisas diferentes. Onde a planilha tem fórmula, o PDF mostra o valor que o sistema calculou.
Sem custo e sem depender do Excel instalado."""
import html as H
import re

from openpyxl.utils import get_column_letter, range_boundaries

PX_POR_CARACTERE, PX_FIXO = 7, 5        # largura de coluna da planilha (em caracteres) -> pixels
MM_POR_POLEGADA, PX_POR_MM = 25.4, 96 / 25.4
A4 = (210, 297)
ESPESSURA = {'thin': '1px solid #000', 'medium': '2px solid #000', 'thick': '3px solid #000', 'hair': '1px solid #888', 'dotted': '1px dotted #000', 'dashed': '1px dashed #000'}


def _area(ws):
    """(coluna inicial, linha inicial, coluna final, linha final) da área de impressão da aba (ou a aba toda)."""
    pa = ws.print_area
    pa = pa[0] if isinstance(pa, (list, tuple)) and pa else pa
    if pa:
        return range_boundaries(str(pa).split('!')[-1].replace('$', ''))
    return range_boundaries(ws.dimensions)


def _cor(c):
    """Cor em #RRGGBB quando a planilha a guarda como RGB (cor de tema: fica a cor padrão)."""
    if c is None or getattr(c, 'type', None) != 'rgb' or not isinstance(c.rgb, str) or len(c.rgb) < 6:
        return None
    return '#' + c.rgb[-6:]


def moeda(v):
    """Número no formato contábil da planilha: ('R$', '1.234,56'); zero aparece como traço; negativo com o sinal na frente."""
    if abs(v) < 0.005:
        return 'R$', '-'
    t = f'{abs(v):,.2f}'.replace(',', '#').replace('.', ',').replace('#', '.')
    return ('-R$' if v < 0 else 'R$'), t


def _valor(ws, cel, calculados):
    v = cel.value
    if isinstance(v, str) and v.startswith('='):
        v = calculados.get(ws.title, {}).get(cel.coordinate)
    if v is None:
        return ''
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        if 'R$' in (cel.number_format or ''):
            a, b = moeda(float(v))
            return f'<span class="m"><span>{a}</span><span>{b}</span></span>'
        return H.escape(str(int(v)) if float(v).is_integer() else f'{v:.2f}'.replace('.', ','))
    return H.escape(str(v)).replace('\n', '<br>')


def _estilo(ws, cel, ultima):
    """CSS da célula. ultima: a célula do canto oposto da mesclagem (de onde vêm as bordas direita e de baixo)."""
    f, al = cel.font, cel.alignment
    css = [f"font-family:'{f.name or 'Arial'}','Arial Narrow',Arial,sans-serif", f'font-size:{f.sz or 11}pt']
    if f.b:
        css.append('font-weight:bold')
    cor = _cor(f.color)
    if cor:
        css.append(f'color:{cor}')
    if cel.fill is not None and cel.fill.fill_type == 'solid':
        fundo = _cor(cel.fill.fgColor)
        if fundo:
            css.append(f'background:{fundo}')
    h = al.horizontal if al.horizontal in ('center', 'left', 'right') else ('right' if isinstance(cel.value, (int, float)) and 'R$' not in (cel.number_format or '') else 'left')
    css.append(f'text-align:{h}')
    css.append(f"vertical-align:{ {'center': 'middle', 'top': 'top'}.get(al.vertical, 'bottom') }")
    css.append('white-space:normal' if al.wrap_text else 'white-space:nowrap')
    for lado, de in (('left', cel), ('top', cel), ('right', ultima), ('bottom', ultima)):
        b = ESPESSURA.get(getattr(getattr(de.border, lado), 'style', None))
        if b:
            css.append(f'border-{lado}:{b}')
    return ';'.join(css)


def html_da_aba(ws, calculados=None):
    """(HTML da tabela, largura em pixels, linhas de título que se repetem em cada página)."""
    calculados = calculados or {}
    c0, l0, c1, l1 = _area(ws)
    larguras = [round((ws.column_dimensions[get_column_letter(c)].width or 8.43) * PX_POR_CARACTERE + PX_FIXO) for c in range(c0, c1 + 1)]
    juntas, cobertas = {}, set()
    for m in ws.merged_cells.ranges:
        a, b, c, d = m.min_col, m.min_row, m.max_col, m.max_row
        if c < c0 or a > c1 or d < l0 or b > l1:
            continue
        juntas[(b, a)] = (min(d, l1), min(c, c1))
        cobertas |= {(l, k) for l in range(b, d + 1) for k in range(a, c + 1)} - {(b, a)}
    titulos = 0
    m = re.match(r'\$?(\d+):\$?(\d+)$', ws.print_title_rows or '')
    if m and int(m.group(1)) == l0:
        titulos = int(m.group(2)) - l0 + 1
    out = ['<table><colgroup>' + ''.join(f'<col style="width:{w}px">' for w in larguras) + '</colgroup>']
    for l in range(l0, l1 + 1):
        if l == l0:
            out.append('<thead>' if titulos else '<tbody>')
        elif titulos and l == l0 + titulos:
            out.append('</thead><tbody>')
        altura = ws.row_dimensions[l].height
        out.append(f'<tr style="height:{altura or 15}pt">')
        for c in range(c0, c1 + 1):
            if (l, c) in cobertas:
                continue
            cel = ws.cell(row=l, column=c)
            fl, fc = juntas.get((l, c), (l, c))
            span = (f' colspan="{fc - c + 1}"' if fc > c else '') + (f' rowspan="{fl - l + 1}"' if fl > l else '')
            if fl > l:   # célula que junta as linhas de uma rubrica: pode continuar na página seguinte (as linhas comuns não são cortadas)
                span += ' class="g"'
            out.append(f'<td{span} style="{_estilo(ws, cel, ws.cell(row=fl, column=fc))}">{_valor(ws, cel, calculados)}</td>')
        out.append('</tr>')
    out.append('</tbody></table>')
    return ''.join(out), sum(larguras), titulos


CSS = """*{box-sizing:border-box;-webkit-print-color-adjust:exact;print-color-adjust:exact}
body{margin:0;background:#fff;color:#000}
table{border-collapse:collapse;table-layout:fixed}
td{padding:1px 3px;overflow:hidden;line-height:1.15}
thead{display:table-header-group}
td{break-inside:avoid}
td.g{break-inside:auto}
.m{display:flex;justify-content:space-between;gap:4px;white-space:nowrap}"""


def paginas_html(wb):
    """Uma página HTML por aba, com o que o navegador precisa para imprimir como a planilha: [dict(titulo, html, paisagem, escala, margens em mm)]."""
    calculados = getattr(wb, '_calculados', {})
    out = []
    for ws in wb.worksheets:
        tabela, largura, _ = html_da_aba(ws, calculados)
        paisagem = ws.page_setup.orientation == 'landscape'
        mg = ws.page_margins
        margens = {k: round((getattr(mg, k) or 0.5) * MM_POR_POLEGADA, 1) for k in ('left', 'right', 'top', 'bottom')}
        util = ((A4[1] if paisagem else A4[0]) - margens['left'] - margens['right']) * PX_POR_MM
        escala = (ws.page_setup.scale or 100) / 100
        cabe = util / largura if largura else 1
        ajustar = bool(ws.sheet_properties.pageSetUpPr and ws.sheet_properties.pageSetUpPr.fitToPage)
        escala = min(1.0, cabe) if ajustar else min(escala, cabe)   # nunca cortar colunas: se nem na escala da planilha couber, diminui
        out.append(dict(titulo=ws.title, paisagem=paisagem, escala=round(max(0.1, escala), 3), margens=margens,
                        html=f'<!doctype html><html lang="pt-BR"><head><meta charset="utf-8"><title>{H.escape(ws.title)}</title><style>{CSS}</style></head><body>{tabela}</body></html>'))
    return out


async def _imprimir(paginas):
    import pymupdf
    from playwright.async_api import async_playwright
    final = pymupdf.open()
    async with async_playwright() as pw:
        br = await pw.chromium.launch(headless=True)
        try:
            pg = await br.new_page()
            for x in paginas:
                await pg.set_content(x['html'])
                pdf = await pg.pdf(format='A4', landscape=x['paisagem'], scale=x['escala'], print_background=True,
                                   margin={k: f'{v}mm' for k, v in x['margens'].items()})
                with pymupdf.open(stream=pdf, filetype='pdf') as parte:
                    final.insert_pdf(parte)
        finally:
            await br.close()
    try:
        return final.tobytes(garbage=3, deflate=True)
    finally:
        final.close()


def pdf_da_planilha(wb):
    """O PDF (bytes) com todas as abas da planilha, na ordem. Abre o navegador do sistema; chamar fora de um laço de eventos em execução."""
    import asyncio
    return asyncio.run(_imprimir(paginas_html(wb)))
