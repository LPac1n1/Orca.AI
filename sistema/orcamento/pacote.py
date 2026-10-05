"""Pacote para a Secretaria (pedido da OSC, 05/10/2026): tudo pronto para enviar, num arquivo .zip.

- "Plano de Aplicação e Comparativo de Preço.xlsx": as duas planilhas com a MESMA formatação da planilha de pré-cálculos da OSC
  (docs-base/Pré-Calculos.xlsx): fontes, bordas, mesclagens, larguras, formato de moeda e configuração de impressão. Médias e totais são fórmulas.
- "Orçamentos/NN. Rubrica/…": o PDF de cada pesquisa (vaga, produto, cotação, proposta), na ordem do plano; cada PDF já traz, no fim, o
  Comprovante de Inscrição e de Situação Cadastral (CNPJ ativo) da empresa daquela pesquisa.
- "LEIA-ME.txt": o que foi incluído e o que ainda falta (pesquisa sem PDF, empresa sem comprovante de CNPJ).
Nada é inventado: pesquisa sem PDF não gera arquivo, e a falta fica escrita no LEIA-ME."""
import datetime as dt
import math
import os
import re
import zipfile

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.page import PageMargins

from . import db
from .exportar import _desc_rh, _qtd_rh, _fornecedor
from .modelo import Projeto, RubricaRH, fontes_do_subitem, descricao_completa
from .regras import brl, cnpj_formatar, media

MOEDA = '_-"R$ "* #,##0.00_-;"-R$ "* #,##0.00_-;_-"R$ "* \\-??_-;_-@_-'
M, F_ = Side(style='medium'), Side(style='thin')
PRETO, BRANCO = PatternFill('solid', fgColor='FF000000'), PatternFill('solid', fgColor='FFFFFFFF')
CENTRO = Alignment(horizontal='center', vertical='center', wrap_text=True)


def _c(ws, ref, valor, fonte, borda=None, fmt=None, fundo=None, alinhamento=CENTRO, calc=None):
    """calc: o resultado da fórmula, em reais (é gravado no arquivo junto com ela — ver _gravar_valores)."""
    c = ws[ref]
    if calc is not None:
        ws.parent.__dict__.setdefault('_calculados', {}).setdefault(ws.title, {})[ref] = round(calc, 2)
    if type(c).__name__ == 'MergedCell':   # parte de uma célula mesclada: só a borda (o valor fica na primeira célula)
        if borda:
            c.border = borda
        return c
    c.value, c.font, c.alignment = valor, fonte, alinhamento
    if borda:
        c.border = borda
    if fmt:
        c.number_format = fmt
    if fundo:
        c.fill = fundo
    return c


def _linhas(texto, por_linha):
    return max(1, sum(math.ceil(len(t) / por_linha) or 1 for t in str(texto or '').split('\n')))


# ------------------------------------------------------------------ Comparativo de Preço
def comparativo(wb, p: Projeto):
    ws = wb.create_sheet('Comparativo de Preço')
    VB, T, TB = Font(name='Verdana', size=8, bold=True), Font(name='Times New Roman', size=8), Font(name='Times New Roman', size=8, bold=True)
    tudo = Border(left=M, right=M, top=M, bottom=M)
    for col in 'ABCDEFGHIJKLM':
        ws.column_dimensions[col].width = 13.3
    ws.merge_cells('A1:M1')
    _c(ws, 'A1', 'COMPARATIVO DE PREÇOS', Font(name='Verdana', size=8, bold=True, color='FFFFFFFF'), tudo, fundo=PRETO)
    for faixa, ref, txt in (('A2:A4', 'A2', 'ITEM'), ('B2:B4', 'B2', 'DESCRIÇÃO DO ITEM'), ('C2:C4', 'C2', 'Q\nT\nD\nE'), ('M2:M4', 'M2', 'Valor Médio'),
                            ('D2:F2', 'D2', 'ORÇAMENTO 1'), ('G2:I2', 'G2', 'ORÇAMENTO 2'), ('J2:L2', 'J2', 'ORÇAMENTO 3')):
        ws.merge_cells(faixa); _c(ws, ref, txt, VB, tudo)
    for a, b, f in (('D', 'E', 'F'), ('G', 'H', 'I'), ('J', 'K', 'L')):
        ws.merge_cells(f'{a}3:{a}4'); _c(ws, f'{a}3', 'PREÇO UNITÁRIO', VB, tudo, MOEDA, BRANCO)
        ws.merge_cells(f'{b}3:{b}4'); _c(ws, f'{b}3', 'PREÇO TOTAL ', VB, tudo, MOEDA, BRANCO)
        _c(ws, f'{f}3', 'CNPJ', VB, Border(right=M, top=M), fundo=BRANCO); _c(ws, f'{f}4', 'FORNECEDOR', VB, Border(right=M, bottom=M), fundo=BRANCO)
    for lin in range(1, 5):
        ws.row_dimensions[lin].height = 15
        for col in 'ABCDEFGHIJKLM':   # as bordas das células mescladas
            c = ws[f'{col}{lin}']
            if not getattr(c.border.left, 'style', None) and not getattr(c.border.right, 'style', None):
                c.border = tudo
    trios = (('D', 'E', 'F'), ('G', 'H', 'I'), ('J', 'K', 'L'))

    def linha(lin, item, desc, qtd, precos, fornecedores, fonte, fina=False, ultima=False, somas=None):
        """precos: os 3 preços unitários em CENTAVOS (linha comum) ou as 3 fórmulas de soma (linha da rubrica; somas = [(unitário, total)] em centavos)."""
        lado = F_ if fina else M
        baixo = M if (not fina or ultima) else F_
        b = Border(right=lado, bottom=baixo)
        _c(ws, f'A{lin}', item, VB, Border(left=M, right=M, bottom=baixo), fundo=BRANCO)
        _c(ws, f'B{lin}', desc, fonte, b); _c(ws, f'C{lin}', qtd, fonte, b)
        for k, (a, t, f) in enumerate(trios):
            fundo = BRANCO if k == 0 else None
            if somas:
                _c(ws, f'{a}{lin}', precos[k], fonte, b, MOEDA, fundo, calc=somas[k][0] / 100)
                _c(ws, f'{t}{lin}', precos[k], fonte, b, MOEDA, fundo, calc=somas[k][1] / 100)
            else:
                _c(ws, f'{a}{lin}', precos[k] / 100, fonte, b, MOEDA, fundo)
                _c(ws, f'{t}{lin}', f'={a}{lin}*C{lin}', fonte, b, MOEDA, fundo, calc=precos[k] * qtd / 100)
            _c(ws, f'{f}{lin}', fornecedores[k], fonte, Border(left=lado, right=lado, bottom=baixo), fundo=fundo)
        ws.row_dimensions[lin].height = 10.5 + 8.5 * max(4, _linhas(desc, 15), *[_linhas(x, 15) for x in fornecedores])

    lin = 5
    for r in p.rubricas:
        if isinstance(r, RubricaRH):
            ps = (list(r.pesquisas) + [None] * 3)[:3]
            val = [((q.valor if q.valor is not None else q.faixa_min) or 0) if q else 0 for q in ps]
            linha(lin, r.item, _desc_rh(r), _qtd_rh(r), val, [_fornecedor(q) if q and q.nome else '' for q in ps], T)
            _c(ws, f'M{lin}', f'=ROUND((D{lin}+G{lin}+J{lin})/3,2)', T, Border(right=M, bottom=M), MOEDA, calc=media(val) / 100)
            lin += 1
            continue
        cab, primeira = lin, lin + 1
        somas = [[0, 0], [0, 0], [0, 0]]   # de cada orçamento: soma dos unitários e soma dos totais (centavos)
        for i, s in enumerate(r.subitens):
            lin += 1
            fs = fontes_do_subitem(r, s)
            val = [(s.precos[k] or 0) if k < len(s.precos) else 0 for k in range(3)]
            for k in range(3):
                somas[k][0] += val[k]; somas[k][1] += val[k] * s.qtd
            linha(lin, None, descricao_completa(s), s.qtd, val,
                  [_fornecedor(fs[k]) if k < len(fs) and fs[k].nome else '' for k in range(3)], T, fina=True, ultima=i == len(r.subitens) - 1)
            _c(ws, f'M{lin}', f'=ROUND((D{lin}+G{lin}+J{lin})/3,2)', T, Border(left=F_, right=M, bottom=M if i == len(r.subitens) - 1 else F_), MOEDA, calc=media(val) / 100)
        ult = lin
        por_item = r.subitens and all(len(s.fontes) == 3 for s in r.subitens)
        nomes = [fontes_do_subitem(r, s) for s in r.subitens]
        mesmos = [(_fornecedor(nomes[0][k]) if nomes and k < len(nomes[0]) and len({(f[k].cnpj if k < len(f) else None) for f in nomes}) == 1 else
                   ('fornecedores por item (abaixo)' if por_item else (_fornecedor(r.fontes[k]) if k < len(r.fontes) and r.fontes[k].nome else ''))) for k in range(3)]
        if ult >= primeira:
            linha(cab, r.item, r.descricao, '', [f'=SUM({a}{primeira}:{a}{ult})' for a, _, _ in trios], mesmos, TB, somas=somas)
            for a, t, _ in trios:
                ws[f'{t}{cab}'].value = f'=SUM({t}{primeira}:{t}{ult})'
            _c(ws, f'M{cab}', f'=ROUND((E{cab}+H{cab}+K{cab})/3,2)', TB, Border(right=M, bottom=M), MOEDA, calc=media([x[1] for x in somas]) / 100)
            ws.merge_cells(f'A{cab}:A{ult}')
            for x in range(cab, ult + 1):
                ws[f'A{x}'].border = Border(left=M, right=M, bottom=M if x == ult else None)
        else:   # rubrica ainda sem itens
            linha(cab, r.item, r.descricao, 0, [0, 0, 0], ['', '', ''], TB)
            _c(ws, f'M{cab}', 0, TB, Border(right=M, bottom=M), MOEDA)
        lin = ult + 1
    ws.print_area = f'A1:M{max(lin - 1, 4)}'
    ws.print_title_rows = '1:4'
    ws.page_setup.orientation, ws.page_setup.paperSize, ws.page_setup.scale = 'landscape', 9, 74
    ws.page_margins = PageMargins(left=0.236111111111111, right=0.236111111111111, top=0.747916666666667, bottom=0.747916666666667,
                                  header=0.511811023622047, footer=0.511811023622047)
    ws.sheet_view.zoomScale = 110
    return ws


# ------------------------------------------------------------------ Plano de Aplicação
def plano(wb, p: Projeto):
    ws = wb.active; ws.title = 'Plano de Aplicação'
    N, B = Font(name='Aptos Narrow', size=9), Font(name='Aptos Narrow', size=9, bold=True)
    esq = Alignment(vertical='center', wrap_text=True)
    meio = Alignment(horizontal='center', vertical='center')
    for col, w in zip('ABCDEFGH', (12.9, 8.3, 4.7, 40.3, 12.0, 12.9, 13.0, 10.1)):
        ws.column_dimensions[col].width = w
    ws.merge_cells('C1:H1')
    _c(ws, 'C1', 'PLANO DE APLICAÇÃO', Font(name='Aptos Narrow', size=9, bold=True, color='FFFFFFFF'), Border(left=F_, right=F_, top=F_), fundo=PRETO, alinhamento=meio)
    for col in 'DEFGH':
        ws[f'{col}1'].border = Border(top=F_, right=F_ if col == 'H' else None); ws[f'{col}1'].fill = PRETO
    ws.row_dimensions[1].height, ws.row_dimensions[2].height = 12, 23

    def faixa(lin, valores, fontes, cima=M, baixo=F_, formatos=(None, None, MOEDA, MOEDA, MOEDA, None), calc=(None,) * 6):
        """calc: o resultado (em centavos) das células que são fórmula."""
        for i, col in enumerate('CDEFGH'):
            borda = Border(left=M if col == 'C' else F_, right=M if col == 'H' else F_, top=cima, bottom=baixo)
            _c(ws, f'{col}{lin}', valores[i], fontes[i], borda, formatos[i], calc=None if calc[i] is None else calc[i] / 100, alinhamento=esq if col == 'D' and lin > 3 else Alignment(horizontal='center', vertical='center', wrap_text=col in 'EGH' or lin == 2))
    faixa(2, ['Item', 'Descrição', 'Valor Unitário', 'Valor Total', 'Concedente\n (SJC)', 'Proponente\n (entidade)'], [B] * 6, M, M)
    lin, totais, geral = 3, [], 0
    rh = [r for r in p.rubricas if isinstance(r, RubricaRH)]
    if rh:
        a, b = 4, 3 + len(rh)
        mensal = sum((r.valor_mensal_plano or 0) * _qtd_rh(r) for r in rh)
        total = sum((r.valor_mensal_plano or 0) * _qtd_rh(r) * r.meses for r in rh)
        faixa(3, [None, 'Recursos Humanos', f'=SUM(E{a}:E{b})', f'=SUM(F{a}:F{b})', f'=SUM(G{a}:G{b})', None], [B] * 6, M, M, calc=(None, None, mensal, total, total, None))
        for col in 'CDEFGH':
            ws[f'{col}3'].fill = BRANCO
        ws.row_dimensions[3].height = 12.5
        lin = 4
    for r in p.rubricas:
        if isinstance(r, RubricaRH):
            q, v = _qtd_rh(r), (r.valor_mensal_plano or 0)
            desc = f'{_desc_rh(r)} ({q} profissionais x {brl(v).replace(" ", "")})' if q > 1 else _desc_rh(r)
            faixa(lin, [r.item, desc, f'={q}*{v / 100}' if q > 1 else v / 100, f'=ROUND(E{lin}*{r.meses},2)', f'=F{lin}', None], [B, N, B, N, N, N],
                  calc=(None, None, q * v if q > 1 else None, q * v * r.meses, q * v * r.meses, None))
            ws.row_dimensions[lin].height = 12.5 + 10.5 * (_linhas(desc, 52) - 1)
            totais.append(lin); geral += q * v * r.meses; lin += 1
            continue
        cab, mensal = lin, 0
        for s in r.subitens:
            lin += 1
            v = s.valor_plano or 0
            mensal += s.qtd * v
            desc = f'{descricao_completa(s)} ({s.qtd} unidade{"s" if s.qtd > 1 else ""} x {brl(v).replace(" ", "")})'
            faixa(lin, [None, desc, f'={s.qtd}*{v / 100}', None, None, None], [N] * 6, F_, F_, calc=(None, None, s.qtd * v, None, None, None))
            ws.row_dimensions[lin].height = 12.5 + 10.5 * (_linhas(desc, 52) - 1)
        ult = lin
        soma = f'=SUM(E{cab + 1}:E{ult})' if ult > cab else 0
        faixa(cab, [r.item, r.descricao, soma, f'=ROUND(E{cab}*{r.meses},2)', f'=F{cab}', None], [B] * 6, M, F_,
              calc=(None, None, mensal if ult > cab else None, mensal * r.meses, mensal * r.meses, None))
        geral += mensal * r.meses
        ws.row_dimensions[cab].height = 12.5 + 10.5 * (_linhas(r.descricao, 48) - 1)
        if ult > cab:
            for col in 'CFG':   # o número do item e os totais valem para a rubrica inteira
                ws.merge_cells(f'{col}{cab}:{col}{ult}')
        totais.append(cab); lin = ult + 1
    V = Font(name='Aptos Narrow', size=9, bold=True, color='FFFF0000')
    ws.merge_cells(f'C{lin}:E{lin}')
    faixa(lin, ['Total', None, None, '=' + ('+'.join(f'F{x}' for x in totais) or '0'), '=' + ('+'.join(f'G{x}' for x in totais) or '0'), 0], [V] * 6, M, M,
          (None, None, None, MOEDA, MOEDA, MOEDA), calc=(None, None, None, geral, geral, None))
    ws.row_dimensions[lin].height = 12
    ws.print_area = f'C1:H{lin}'
    ws.page_setup.orientation, ws.page_setup.paperSize = 'portrait', 9
    ws.page_margins = PageMargins(left=0.25, right=0.25, top=0.75, bottom=0.75)
    return ws


def _gravar_valores(caminho, wb):
    """Grava no arquivo, junto de cada fórmula, o valor já calculado pelo sistema. Sem isso, a planilha só mostra os números depois que o Excel
    recalcula — e ele NÃO recalcula no "Modo de Exibição Protegido" (como abre todo arquivo baixado pelo navegador), nem a pré-visualização do
    Windows, o celular ou o e-mail: as fórmulas aparecem zeradas e os valores parecem diferentes dos do sistema (relato da OSC, 05/10/2026).
    As fórmulas continuam no arquivo: quem editar um número vê os totais mudarem."""
    calculados = getattr(wb, '_calculados', {})
    por_arquivo = {ws.path.lstrip('/'): calculados.get(ws.title, {}) for ws in wb.worksheets}
    tmp = caminho + '.valores'
    with zipfile.ZipFile(caminho) as zin, zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            dados = zin.read(info.filename)
            vals = por_arquivo.get(info.filename)
            if vals:
                def com_valor(m):
                    v = vals.get(m.group(2))
                    return m.group(1) + (f'<v>{v:.2f}</v>' if v is not None else m.group(3))
                dados = re.sub(r'(<c r="([A-Z]+[0-9]+)"[^>]*><f>[^<]*</f>)(<v\s*/>|<v></v>)', com_valor, dados.decode('utf-8')).encode('utf-8')
            zout.writestr(info, dados)
    os.replace(tmp, caminho)


def planilhas(p: Projeto, caminho):
    """Plano de Aplicação + Comparativo de Preço, na formatação da planilha de pré-cálculos. Cada fórmula leva junto o valor calculado."""
    wb = Workbook()
    plano(wb, p); comparativo(wb, p)
    wb.save(caminho)
    _gravar_valores(caminho, wb)
    return caminho


# ------------------------------------------------------------------ PDFs por rubrica, com o comprovante de CNPJ junto
def _nome(texto, limite=70):
    t = re.sub(r'[\\/:*?"<>|\r\n\t]+', ' ', str(texto or '')).strip(' .')
    return re.sub(r'\s+', ' ', t)[:limite].strip(' .') or 'sem nome'


def _juntar(principal, comprovante):
    """O PDF da pesquisa com o comprovante de CNPJ no fim (um arquivo só). Sem comprovante, só o PDF da pesquisa."""
    if not comprovante:
        with open(principal, 'rb') as fh:
            return fh.read()
    import fitz
    with fitz.open(principal) as doc, fitz.open(comprovante) as extra:
        doc.insert_pdf(extra)
        return doc.tobytes(garbage=3, deflate=True)


def pesquisas_do_projeto(p: Projeto):
    """Cada pesquisa do orçamento, na ordem do plano: dict(pasta, arquivo, empresa, cnpj, pdf (caminho relativo ou None), onde)."""
    out = []
    for r in p.rubricas:
        pasta = f'{r.item:02d}. {_nome(r.cargo if isinstance(r, RubricaRH) else r.descricao, 60)}'
        if isinstance(r, RubricaRH):
            for k, q in enumerate(r.pesquisas[:3]):
                if q.nome or q.evidencia.arquivo:
                    out.append(dict(pasta=pasta, arquivo=f'{k + 1}. {_nome(q.nome or "empresa não informada")}.pdf', empresa=q.nome, cnpj=q.cnpj,
                                    pdf=q.evidencia.arquivo, onde=f'item {r.item} ({r.cargo}), pesquisa {k + 1}'))
            continue
        varios = len(r.subitens) > 1
        for i, s in enumerate(r.subitens):
            sub = f'{pasta}/{i + 1:02d}. {_nome(descricao_completa(s), 60)}' if varios else pasta
            for k, f in enumerate(fontes_do_subitem(r, s)[:3]):
                if f.nome or f.evidencia.arquivo:
                    out.append(dict(pasta=sub, arquivo=f'{k + 1}. {_nome(f.plataforma or f.nome or "fornecedor não informado")}.pdf', empresa=f.nome, cnpj=f.cnpj,
                                    pdf=f.evidencia.arquivo, onde=f'item {r.item} ({r.descricao}) / {descricao_completa(s)}, pesquisa {k + 1}'))
    return out


def montar(p: Projeto, pid, caminho_zip):
    """Grava o .zip e devolve dict(arquivos, com_cnpj, sem_pdf=[...], sem_cnpj=[...])."""
    raiz = _nome(p.nome, 80)
    sem_pdf, sem_cnpj, n, com = [], [], 0, 0
    tmp = caminho_zip + '.xlsx'
    planilhas(p, tmp)
    with zipfile.ZipFile(caminho_zip, 'w', zipfile.ZIP_DEFLATED) as z:
        z.write(tmp, f'{raiz}/Plano de Aplicação e Comparativo de Preço.xlsx')
        usados = set()
        for x in pesquisas_do_projeto(p):
            ab = db.caminho_absoluto(x['pdf']) if x['pdf'] else None
            if not ab or not os.path.exists(ab):
                sem_pdf.append(x['onde'] + (f' — {x["empresa"]}' if x['empresa'] else '')); continue
            comp = p.comprovantes_cnpj.get(cnpj_formatar(x['cnpj'])) if x['cnpj'] else None
            ab_c = db.caminho_absoluto(comp['arquivo']) if comp and comp.get('arquivo') else None
            if not ab_c or not os.path.exists(ab_c):
                ab_c = None
                sem_cnpj.append(f'{x["onde"]} — {x["empresa"] or "empresa não informada"}' + (f' ({cnpj_formatar(x["cnpj"])})' if x['cnpj'] else ' (sem CNPJ)'))
            try:
                dados = _juntar(ab, ab_c)
            except Exception as e:   # PDF que não abre: entra como está, e a falta do comprovante fica registrada
                with open(ab, 'rb') as fh:
                    dados = fh.read()
                if ab_c:
                    sem_cnpj.append(f'{x["onde"]} — não foi possível juntar o comprovante de CNPJ ao PDF ({type(e).__name__})')
                ab_c = None
            destino = f'{raiz}/Orçamentos/{x["pasta"]}/{x["arquivo"]}'
            while destino in usados:   # duas pesquisas com a mesma empresa na mesma pasta
                destino = destino[:-4] + ' (2).pdf'
            usados.add(destino)
            z.writestr(destino, dados)
            n += 1; com += 1 if ab_c else 0
        leia = [f'{p.nome}', f'Pacote gerado em {dt.datetime.now().strftime("%d/%m/%Y %H:%M")}.', '',
                'O QUE HÁ AQUI',
                '- "Plano de Aplicação e Comparativo de Preço.xlsx": as duas planilhas (médias e totais são fórmulas).',
                f'- Pasta "Orçamentos": {n} PDF(s), separados por rubrica e por item, na ordem do plano. Cada PDF é a página da vaga ou do produto',
                f'  (ou a proposta) e, no fim do mesmo arquivo, o Comprovante de Inscrição e de Situação Cadastral da empresa ({com} de {n} já com o comprovante).', '']
        if sem_pdf or sem_cnpj:
            leia.append('O QUE AINDA FALTA')
            leia += [f'- Sem PDF da pesquisa (não entrou no pacote): {t}' for t in sem_pdf]
            leia += [f'- Sem o comprovante de CNPJ junto: {t}' for t in sem_cnpj]
            leia += ['', 'Para completar: tela do projeto → "CNPJs e comprovantes" (emitir e importar os comprovantes) e "Verificação" (pendências das pesquisas).']
        else:
            leia.append('Nada falta: todas as pesquisas têm o PDF e o comprovante de CNPJ da empresa.')
        z.writestr(f'{raiz}/LEIA-ME.txt', '\r\n'.join(leia))
    try:
        os.remove(tmp)
    except OSError:
        pass
    return dict(arquivos=n, com_cnpj=com, sem_pdf=sem_pdf, sem_cnpj=sem_cnpj)
