"""Avaliador das fórmulas que o Orça.AI escreve nas planilhas (só para os testes).

Aceita o que os geradores usam: números, textos, referências (inclusive de outra aba),
intervalos, + - * /, comparações e as funções SUM, AVERAGE, ROUND e IF. As contas são
exatas (Fraction): o teste confere a lógica das fórmulas contra os valores do dossiê.
"""

import io
import re
from decimal import Decimal
from fractions import Fraction

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter

_REF_ABA = re.compile(r"'([^']+)'!\$?([A-Z]{1,3})\$?(\d+)")
_INTERVALO = re.compile(r"\$?([A-Z]{1,3})\$?(\d+):\$?([A-Z]{1,3})\$?(\d+)")
_REF = re.compile(r"(?<![A-Za-z_'\"])\$?([A-Z]{1,3})\$?(\d+)(?![\w(])")


def _round(valor, casas):
    fator = Fraction(10) ** int(casas)
    x = Fraction(valor) * fator
    sinal = -1 if x < 0 else 1
    return sinal * ((abs(x) + Fraction(1, 2)).__floor__()) / fator


def _achatar(args):
    for a in args:
        if isinstance(a, list):
            yield from _achatar(a)
        else:
            yield a


class Planilha:
    def __init__(self, conteudo: bytes):
        self.livro = load_workbook(io.BytesIO(conteudo))
        self._cache: dict[tuple[str, str], object] = {}

    def valor(self, aba: str, celula: str):
        chave = (aba, celula)
        if chave not in self._cache:
            bruto = self.livro[aba][celula].value
            self._cache[chave] = self._avaliar(aba, bruto)
        return self._cache[chave]

    def _numero(self, v):
        if isinstance(v, bool) or v is None:
            return v if v is not None else Fraction(0)
        if isinstance(v, (int, Decimal, float)):
            return Fraction(v) if not isinstance(v, float) else Fraction(str(v))
        return v

    def _avaliar(self, aba: str, bruto):
        if not (isinstance(bruto, str) and bruto.startswith("=")):
            return self._numero(bruto)
        formula = bruto[1:]
        partes = re.split(r'("[^"]*")', formula)
        convertidas = []
        for parte in partes:
            if parte.startswith('"'):
                convertidas.append(parte)
                continue
            parte = _REF_ABA.sub(lambda m: f"_v({m.group(1)!r},'{m.group(2)}{m.group(3)}')", parte)
            parte = _INTERVALO.sub(lambda m: f"_intervalo(_aba,'{m.group(1)}',{m.group(2)},'{m.group(3)}',{m.group(4)})", parte)
            parte = _REF.sub(lambda m: f"_v(_aba,'{m.group(1)}{m.group(2)}')", parte)
            parte = parte.replace("<>", "!=")
            parte = re.sub(r"(?<![<>!=])=(?!=)", "==", parte)
            for excel, py in (("ROUND(", "_round("), ("AVERAGE(", "_media("), ("SUM(", "_soma("), ("IF(", "_se(")):
                parte = parte.replace(excel, py)
            convertidas.append(parte)
        expressao = "".join(convertidas)
        ambiente = {
            "_aba": aba,
            "_v": lambda a, c: self.valor(a, c),
            "_intervalo": self._intervalo,
            "_round": _round,
            "_media": lambda *a: (lambda v: sum(v, Fraction(0)) / len(v))(list(_achatar(a))),
            "_soma": lambda *a: sum(_achatar(a), Fraction(0)),
            "_se": lambda c, s, n: s if c else n,
        }
        return eval(expressao, {"__builtins__": {}}, ambiente)  # noqa: S307 — só fórmulas geradas pelo próprio sistema

    def _intervalo(self, aba, c1, l1, c2, l2):
        return [
            self.valor(aba, f"{get_column_letter(c)}{l}")
            for l in range(l1, l2 + 1)
            for c in range(column_index_from_string(c1), column_index_from_string(c2) + 1)
        ]


def centavos(valor) -> int:
    x = Fraction(valor) * 100
    assert x.denominator == 1, f"valor com mais de 2 casas: {valor}"
    return int(x)
