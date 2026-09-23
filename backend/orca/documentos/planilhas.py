"""Planilhas Excel com fórmulas vivas (docs/02 §18): grade comparativa e plano de aplicação.

As fórmulas refazem as contas na própria planilha, e o teste confere que elas dão
exatamente os valores do dossiê. Médias e totais usam ARRED(…; 2) (ROUND): como a
soma de centavos dividida por 3 nunca termina em meio centavo, o arredondamento
do Excel coincide com o arredondamento comercial do sistema (D-21).
"""

import io
import re
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from orca.documentos.modelo import Dossie, LoteDoc, OrcamentoDoc
from orca.dominio import formatar_cnpj

MOEDA = '"R$" #,##0.00'
NEGRITO = Font(bold=True)
TITULO = Font(bold=True, size=13)
CABECALHO = PatternFill("solid", fgColor="DDE4EE")
DESTAQUE = PatternFill("solid", fgColor="F3F5F8")
FINO = Side(style="thin", color="A0A7B4")
BORDA = Border(left=FINO, right=FINO, top=FINO, bottom=FINO)
CENTRO = Alignment(horizontal="center", vertical="center", wrap_text=True)
REGIMES = {"mei": "MEI", "recibo": "Recibo (RPA)", "clt": "CLT"}


def reais(centavos: int) -> Decimal:
    return Decimal(centavos) / 100


def nome_de_aba(nome: str, usados: set[str]) -> str:
    base = re.sub(r"[\[\]:*?/\\]", "-", nome).strip()[:28] or "Aba"
    candidato, n = base, 2
    while candidato.lower() in usados:
        candidato, n = f"{base[:25]} ({n})", n + 1
    usados.add(candidato.lower())
    return candidato


def _salvar(livro: Workbook) -> bytes:
    saida = io.BytesIO()
    livro.save(saida)
    return saida.getvalue()


def _cabecalho(celula, texto: str) -> None:
    celula.value = texto
    celula.font = NEGRITO
    celula.fill = CABECALHO
    celula.alignment = CENTRO
    celula.border = BORDA


def _cnpj(cnpj: str | None) -> str:
    return f"CNPJ {formatar_cnpj(cnpj)}" if cnpj else "CNPJ não identificado"


# --- Grade comparativa -------------------------------------------------------------------------


def _aba_do_lote(aba, lote: LoteDoc, orcamento: OrcamentoDoc, dossie: Dossie) -> None:
    n = len(lote.lojas)
    col_preco = [4 + 2 * k for k in range(n)]
    col_total = [c + 1 for c in col_preco]
    c_media, c_media_total, c_final, c_total_final, c_situacao = (4 + 2 * n + i for i in range(5))
    L = get_column_letter

    aba["A1"] = f"Grade comparativa — {orcamento.nome} — lote {lote.nome}"
    aba["A1"].font = TITULO
    aba["A2"] = f"Projeto: {dossie.projeto.nome} · {dossie.projeto.organizacao.nome}"
    regra = (
        "B: preço real da loja de menor total (Orçamento 1), que precisa estar dentro da média"
        if lote.base_preco_final == "B" else "A: média dos orçamentos, arredondada"
    )
    aba["A3"] = f"Preço final — regra {regra}. Quantidades finais do período (D-25)."

    for col, texto in ((1, "Item"), (2, "Unid."), (3, "Qtd.")):
        aba.merge_cells(start_row=5, start_column=col, end_row=8, end_column=col)
        _cabecalho(aba.cell(5, col), texto)
    for k, loja in enumerate(lote.lojas):
        c = col_preco[k]
        for linha, texto in ((5, f"Orçamento {k + 1}"), (6, loja.empresa.nome), (7, _cnpj(loja.empresa.cnpj))):
            aba.merge_cells(start_row=linha, start_column=c, end_row=linha, end_column=c + 1)
            _cabecalho(aba.cell(linha, c), texto)
        _cabecalho(aba.cell(8, c), "Preço unit.")
        _cabecalho(aba.cell(8, c + 1), "Total")
    for col, texto in ((c_media, "Média unitária"), (c_media_total, "Média do total"), (c_final, "Preço final"),
                       (c_total_final, "Total final"), (c_situacao, "Situação")):
        aba.merge_cells(start_row=5, start_column=col, end_row=8, end_column=col)
        _cabecalho(aba.cell(5, col), texto)

    primeira = 9
    for i, linha in enumerate(lote.linhas):
        r = primeira + i
        aba.cell(r, 1, linha.nome)
        aba.cell(r, 2, linha.unidade)
        aba.cell(r, 3, linha.quantidade)
        precos = [f"{L(c)}{r}" for c in col_preco]
        for k, fonte in enumerate(linha.fontes):
            aba.cell(r, col_preco[k], reais(fonte.preco_centavos))
            aba.cell(r, col_total[k], f"=ROUND({L(col_preco[k])}{r}*$C{r},2)")
        aba.cell(r, c_media, f"=ROUND(AVERAGE({','.join(precos)}),2)")
        aba.cell(r, c_media_total, f"=ROUND(AVERAGE({','.join(f'{L(c)}{r}' for c in col_total)}),2)")
        if lote.base_preco_final == "B" and linha.preco_final_centavos == linha.fontes[0].preco_centavos:
            aba.cell(r, c_final, f"={precos[0]}")
        elif lote.base_preco_final == "A":
            aba.cell(r, c_final, f"={L(c_media)}{r}")
        else:
            aba.cell(r, c_final, reais(linha.preco_final_centavos))
        aba.cell(r, c_total_final, f"=ROUND({L(c_final)}{r}*$C{r},2)")
        if lote.base_preco_final == "B":
            final = f"{L(c_final)}{r}"
            if linha.comparar_com == "media_exata":
                condicao = f"ROUND({n}*{final},2)<=ROUND({'+'.join(precos)},2)"
            else:
                condicao = f"{final}<={L(c_media)}{r}"
            aba.cell(r, c_situacao, f'=IF({condicao},"dentro da média","ACIMA DA MÉDIA")')
        else:
            aba.cell(r, c_situacao, "média (regra A)")
        for c in range(1, c_situacao + 1):
            aba.cell(r, c).border = BORDA
            if c >= 4 and c not in (c_situacao,):
                aba.cell(r, c).number_format = MOEDA
    ultima = primeira + len(lote.linhas) - 1
    r = ultima + 1
    aba.cell(r, 1, "Total").font = NEGRITO
    for c in (*col_total, c_total_final):
        aba.cell(r, c, f"=SUM({L(c)}{primeira}:{L(c)}{ultima})")
        aba.cell(r, c).number_format = MOEDA
        aba.cell(r, c).font = NEGRITO
    for c in range(1, c_situacao + 1):
        aba.cell(r, c).fill = DESTAQUE
        aba.cell(r, c).border = BORDA
    aba.cell(r + 2, 1, "Orçamentos em ordem de total (Orçamento 1 = menor total, D-26). "
                       "Fontes, datas e evidências: pasta 02_COTACOES do pacote.")
    aba.column_dimensions["A"].width = 38
    aba.column_dimensions["B"].width = 7
    aba.column_dimensions["C"].width = 7
    for c in range(4, c_situacao + 1):
        aba.column_dimensions[L(c)].width = 14
    aba.column_dimensions[L(c_situacao)].width = 18
    aba.freeze_panes = "D9"


def grade_comparativa(orcamento: OrcamentoDoc, dossie: Dossie) -> bytes:
    """Uma aba por lote da rubrica."""
    livro = Workbook()
    livro.remove(livro.active)
    usados: set[str] = set()
    for lote in orcamento.lotes:
        _aba_do_lote(livro.create_sheet(nome_de_aba(lote.nome, usados)), lote, orcamento, dossie)
    return _salvar(livro)


# --- Plano de aplicação e cronogramas -----------------------------------------------------------


def _linhas_do_plano(dossie: Dossie):
    """(orçamento, descrição, unidade, valor unitário, quantidade, meses, mês inicial) na ordem do plano."""
    for orcamento in dossie.orcamentos:
        linhas = []
        for lote in orcamento.lotes:
            for l in lote.linhas:
                linhas.append((l.nome, l.unidade, l.preco_final_centavos, l.quantidade, l.meses, l.mes_inicio))
        for c in orcamento.cargos:
            linhas.append((f"{c.nome} ({REGIMES.get(c.regime, c.regime)})", "posto/mês", c.valor_mensal_centavos, c.postos, c.meses, c.mes_inicio))
        yield orcamento, linhas


def plano_de_aplicacao(dossie: Dossie) -> bytes:
    p = dossie.projeto
    livro = Workbook()
    recursos = livro.active
    recursos.title = "Recursos"
    aplicacao = livro.create_sheet("Aplicação")
    cronograma = livro.create_sheet("Cronograma físico-financeiro")
    desembolso = livro.create_sheet("Cronograma de desembolso")
    L = get_column_letter

    # Aplicação
    aplicacao["A1"] = f"Plano de aplicação — {p.nome}"
    aplicacao["A1"].font = TITULO
    cabecalhos = ["Orçamento", "Item / cargo", "Unid.", "Valor unitário", "Quantidade", "Meses", "Total",
                  "Mês inicial", "Mês final"]
    for c, texto in enumerate(cabecalhos, 1):
        _cabecalho(aplicacao.cell(3, c), texto)
    r = 4
    linhas_plano: list[tuple[int, int, int]] = []  # (linha na aba, mês inicial, meses)
    subtotais = []
    for orcamento, linhas in _linhas_do_plano(dossie):
        inicio = r
        for nome, unidade, valor, qtd, meses, mes_inicio in linhas:
            valores = (orcamento.nome, nome, unidade, reais(valor), qtd, meses, f"=ROUND(D{r}*E{r}*F{r},2)",
                       mes_inicio, mes_inicio + meses - 1)
            for c, v in enumerate(valores, 1):
                aplicacao.cell(r, c, v).border = BORDA
            aplicacao.cell(r, 4).number_format = MOEDA
            aplicacao.cell(r, 7).number_format = MOEDA
            linhas_plano.append((r, mes_inicio, meses))
            r += 1
        aplicacao.cell(r, 1, orcamento.nome)
        aplicacao.cell(r, 2, f"Subtotal — {orcamento.nome}").font = NEGRITO
        aplicacao.cell(r, 7, f"=SUM(G{inicio}:G{r - 1})").number_format = MOEDA
        aplicacao.cell(r, 7).font = NEGRITO
        for c in range(1, 10):
            aplicacao.cell(r, c).fill = DESTAQUE
        subtotais.append(r)
        r += 1
    total = r + 1
    aplicacao.cell(total, 2, "Total do projeto").font = NEGRITO
    aplicacao.cell(total, 7, "=" + "+".join(f"G{s}" for s in subtotais)).number_format = MOEDA
    aplicacao.cell(total, 7).font = NEGRITO
    aplicacao.cell(total + 1, 2, "Teto do projeto")
    aplicacao.cell(total + 1, 7, reais(p.teto_centavos)).number_format = MOEDA
    aplicacao.cell(total + 2, 2, "Diferença para o teto")
    aplicacao.cell(total + 2, 7, f"=G{total + 1}-G{total}").number_format = MOEDA
    for coluna, largura in zip("ABCDEFGHI", (24, 40, 10, 15, 11, 7, 15, 10, 10)):
        aplicacao.column_dimensions[coluna].width = largura
    aplicacao.freeze_panes = "C4"

    # Recursos
    recursos["A1"] = f"Recursos — {p.nome}"
    recursos["A1"].font = TITULO
    recursos["A3"], recursos["B3"] = "Valor total do projeto", f"='Aplicação'!G{total}"
    recursos["A4"], recursos["B4"] = "Contrapartida da OSC", reais(p.contrapartida_centavos)
    recursos["A5"], recursos["B5"] = "Recursos do concedente", "=B3-B4"
    recursos["A6"], recursos["B6"] = "Concedente", p.orgao or "—"
    recursos["A7"], recursos["B7"] = "Instrumento", " ".join(x for x in (p.instrumento, p.processo) if x) or "—"
    for celula in ("B3", "B4", "B5"):
        recursos[celula].number_format = MOEDA
    recursos.column_dimensions["A"].width = 28
    recursos.column_dimensions["B"].width = 40

    # Cronograma físico-financeiro
    meses_projeto = p.duracao_meses
    cronograma["A1"] = f"Cronograma físico-financeiro — {p.nome}"
    cronograma["A1"].font = TITULO
    _cabecalho(cronograma.cell(3, 1), "Orçamento")
    _cabecalho(cronograma.cell(3, 2), "Item / cargo")
    for m in range(1, meses_projeto + 1):
        _cabecalho(cronograma.cell(3, 2 + m), f"Mês {m}")
    c_total = 3 + meses_projeto
    _cabecalho(cronograma.cell(3, c_total), "Total")
    r = 4
    for linha_aplicacao, mes_inicio, meses in linhas_plano:
        cronograma.cell(r, 1, f"='Aplicação'!A{linha_aplicacao}")
        cronograma.cell(r, 2, f"='Aplicação'!B{linha_aplicacao}")
        for m in range(1, meses_projeto + 1):
            ativo = mes_inicio <= m < mes_inicio + meses
            valor = f"=ROUND('Aplicação'!D{linha_aplicacao}*'Aplicação'!E{linha_aplicacao},2)" if ativo else 0
            cronograma.cell(r, 2 + m, valor).number_format = MOEDA
        cronograma.cell(r, c_total, f"=SUM({L(3)}{r}:{L(2 + meses_projeto)}{r})").number_format = MOEDA
        r += 1
    ultima = r - 1
    cronograma.cell(r, 2, "Total do mês").font = NEGRITO
    cronograma.cell(r + 1, 2, "Acumulado").font = NEGRITO
    for m in range(1, meses_projeto + 2):
        col = L(2 + m)
        cronograma.cell(r, 2 + m, f"=SUM({col}4:{col}{ultima})").number_format = MOEDA
        if m <= meses_projeto:
            anterior = f"+{L(1 + m)}{r + 1}" if m > 1 else ""
            cronograma.cell(r + 1, 2 + m, f"={col}{r}{anterior}").number_format = MOEDA
    linha_total_mes = r
    cronograma.column_dimensions["A"].width = 22
    cronograma.column_dimensions["B"].width = 36
    cronograma.freeze_panes = "C4"

    # Cronograma de desembolso (P-06)
    desembolso["A1"] = f"Cronograma de desembolso — {p.nome}"
    desembolso["A1"].font = TITULO
    _cabecalho(desembolso.cell(3, 1), "Mês")
    _cabecalho(desembolso.cell(3, 2), "Parcela")
    for m in range(1, meses_projeto + 1):
        desembolso.cell(3 + m, 1, f"Mês {m}")
        if p.desembolso == "conforme_cronograma":
            valor = f"='Cronograma físico-financeiro'!{L(2 + m)}{linha_total_mes}"
        else:
            valor = f"='Aplicação'!G{total}" if m == 1 else 0
        desembolso.cell(3 + m, 2, valor).number_format = MOEDA
    desembolso.cell(4 + meses_projeto, 1, "Total").font = NEGRITO
    desembolso.cell(4 + meses_projeto, 2, f"=SUM(B4:B{3 + meses_projeto})").number_format = MOEDA
    desembolso["D3"] = (
        "Parcela única no 1º mês (P-06)" if p.desembolso != "conforme_cronograma" else "Conforme o cronograma físico-financeiro"
    )
    desembolso.column_dimensions["A"].width = 12
    desembolso.column_dimensions["B"].width = 16
    return _salvar(livro)
